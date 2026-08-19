"""
M&A Modeling Tool - Excel Generator

Generates the M&A Excel workbook from validated two-company SEC EDGAR data.
Phase 3 tabs: Assumptions, Sources & Uses, Purchase Price Allocation.

Styling conventions carried over from AIO LBO
(/reference/backend/excel_generator.py):
- Blue font = hardcoded input (fetched data + user assumptions)
- Black font = formula
- Yellow fill + comment = defaulted/soft-requirement field
- Red conditional formatting on a nonzero balance check

NON-NEGOTIABLE RULE (AIO LBO's hard-won lesson): every cross-sheet formula
uses a direct 'Sheet Name'!$Cell$Ref — NEVER a named range. A named range
in a cross-sheet formula silently returned $0 in Google Sheets and
LibreOffice and zeroed out an entire tab. Named ranges may be DEFINED for
documentation, but no formula that crosses sheets may use one.

There is deliberately NO debt schedule / cash sweep circularity in this
model: acquisition debt is a fixed amount set at close.
"""

from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName


# =============================================================================
# STYLE CONSTANTS - IB/PE Modeling Convention (carried over from AIO LBO)
# =============================================================================

BLUE_FONT = Font(name='Calibri', size=11, color='0000CC', bold=False)
BLACK_FONT = Font(name='Calibri', size=11, color='000000', bold=False)
BLUE_FONT_BOLD = Font(name='Calibri', size=11, color='0000CC', bold=True)
BLACK_FONT_BOLD = Font(name='Calibri', size=11, color='000000', bold=True)

SECTION_HEADER_FONT = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
SECTION_HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

SUBSECTION_FONT = Font(name='Calibri', size=11, bold=True, color='000000')
SUBSECTION_FILL = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')

DEFAULTED_FILL = PatternFill(start_color='FFFF99', end_color='FFFF99', fill_type='solid')
USER_PROVIDED_FILL = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
ERROR_FILL = PatternFill(start_color='FF6666', end_color='FF6666', fill_type='solid')

# Implausible calculation flag (e.g., COVID-distorted growth rate) - orange,
# distinct from yellow "defaulted" (carried over from AIO LBO)
IMPLAUSIBLE_CALC_FILL = PatternFill(start_color='FFB366', end_color='FFB366', fill_type='solid')

# Headline results styling (Accretion/Dilution summary)
HEADLINE_FONT = Font(name='Calibri', size=16, bold=True, color='000000')
HEADLINE_LABEL_FONT = Font(name='Calibri', size=12, bold=True, color='444444')

RIGHT_ALIGN = Alignment(horizontal='right')

CURRENCY_FORMAT = '_($* #,##0_);_($* (#,##0);_($* "-"??_);_(@_)'
PERCENT_FORMAT = '0.0%'
NUMBER_FORMAT = '#,##0'
PRICE_FORMAT = '$#,##0.00'
EPS_FORMAT = '$#,##0.00'
YEARS_FORMAT = '#,##0'


# =============================================================================
# HELPERS
# =============================================================================

def calculate_revenue_growth_rate(profile: dict) -> tuple:
    """
    Default revenue growth rate from 5-year history using the MEDIAN of YoY
    changes, with a sanity band. Carried over EXACTLY from AIO LBO — the
    original mean-based version produced a 140%+ nonsense default for a
    COVID-distorted company.

    Sanity band: outside -15%..+25% falls back to a conservative 3% with a
    visible flag.

    Returns (growth_rate, is_fallback, comment).
    """
    import statistics

    fiscal_years = profile.get('fiscal_years', {})
    if len(fiscal_years) < 2:
        return 0.03, True, "Insufficient historical data — defaulted to 3%"

    sorted_years = sorted(fiscal_years.keys())
    growth_rates = []

    for i in range(1, len(sorted_years)):
        prev_rev = fiscal_years.get(sorted_years[i - 1], {}).get('revenue')
        curr_rev = fiscal_years.get(sorted_years[i], {}).get('revenue')
        if prev_rev and curr_rev and prev_rev > 0:
            growth_rates.append((curr_rev - prev_rev) / prev_rev)

    if not growth_rates:
        return 0.03, True, "No valid revenue history — defaulted to 3%"

    median_growth = statistics.median(growth_rates)

    GROWTH_MIN = -0.15
    GROWTH_MAX = 0.25
    FALLBACK_RATE = 0.03

    if median_growth < GROWTH_MIN or median_growth > GROWTH_MAX:
        return (
            FALLBACK_RATE,
            True,
            f"Historical growth rate ({median_growth*100:.1f}%) was implausible "
            f"(likely COVID-era distortion) — defaulted to 3%, please review "
            f"and adjust manually."
        )

    return (
        median_growth,
        False,
        f"Default: median of historical YoY growth rates ({median_growth*100:.1f}%)"
    )


def _most_recent_fy_data(profile: dict) -> tuple:
    """Return (fy, fy_data) for the most recent fiscal year, or (None, {})."""
    fiscal_years = profile.get("fiscal_years", {})
    if not fiscal_years:
        return None, {}
    fy = max(fiscal_years.keys())
    return fy, fiscal_years[fy]


def _define_name(wb: Workbook, name: str, sheet_name: str, cell: str):
    """
    Create a workbook-level defined name for DOCUMENTATION ONLY.
    Formulas never use these (see module docstring).
    """
    clean = ''.join(c for c in name.replace(' ', '_').replace('%', 'Pct')
                    if c.isalnum() or c == '_')
    if clean[0].isdigit():
        clean = '_' + clean
    col = ''.join(c for c in cell if c.isalpha())
    row = ''.join(c for c in cell if c.isdigit())
    wb.defined_names[clean] = DefinedName(
        clean, attr_text=f"'{sheet_name}'!${col}${row}")


def _abs_ref(sheet: str, col: str, row: int) -> str:
    """Build a direct cross-sheet reference: 'Sheet'!$C$5"""
    return f"'{sheet}'!${col}${row}"


class RowWriter:
    """Tracks the current row while writing label/value rows to a sheet."""

    def __init__(self, ws, sheet_name: str):
        self.ws = ws
        self.sheet_name = sheet_name
        self.row = 1

    def section(self, title: str, span_end_col: int = 4):
        c = self.ws.cell(row=self.row, column=2, value=title)
        c.font = SECTION_HEADER_FONT
        c.fill = SECTION_HEADER_FILL
        self.ws.merge_cells(start_row=self.row, start_column=2,
                            end_row=self.row, end_column=span_end_col)
        self.row += 2

    def subsection(self, title: str, span_end_col: int = 4):
        c = self.ws.cell(row=self.row, column=2, value=title)
        c.font = SUBSECTION_FONT
        c.fill = SUBSECTION_FILL
        self.ws.merge_cells(start_row=self.row, start_column=2,
                            end_row=self.row, end_column=span_end_col)
        self.row += 1

    def blank(self, n: int = 1):
        self.row += n


def _put_value(ws, row: int, col: int, value, font, number_format=None,
               fill=None, comment: Optional[str] = None):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = font
    if number_format:
        cell.number_format = number_format
    if fill:
        cell.fill = fill
    if comment:
        cell.comment = Comment(comment, "System")
    return cell


# =============================================================================
# TAB 1: ASSUMPTIONS
# =============================================================================

# Column layout: B = label, C = Acquirer, D = Target (E for notes if needed)
ACQ_COL, TGT_COL = 'C', 'D'

# Default deal assumption inputs (user-editable in the sheet; overridable
# programmatically via generate_ma_workbook(deal_assumptions=...))
DEFAULT_DEAL_ASSUMPTIONS = {
    "offer_premium": 0.25,
    "pct_cash": 0.40,
    "pct_debt": 0.30,
    "new_debt_rate": 0.08,
    "foregone_cash_rate": 0.045,
    "revenue_synergies_pct": 0.0,
    "cost_synergies_pct": 0.0,
    "synergy_ramp_years": 3,
    "transaction_fee_pct": 0.02,
    "asset_step_up_pct": 0.50,
    "intangible_amort_years": 10,
    "tax_rate": 0.25,
    "analysis_years": 5,
}


def build_assumptions_tab(wb: Workbook, pair: dict, pair_verdict: dict,
                          deal_assumptions: Optional[dict] = None) -> dict:
    """
    Build the Assumptions tab with Acquirer and Target side by side, plus
    the Deal Assumptions inputs.

    Returns a refs dict mapping semantic names to direct cross-sheet
    references (e.g. refs['tgt_price'] == "'Assumptions'!$D$10") for use
    by every downstream tab.
    """
    ws = wb.create_sheet("Assumptions", 0)
    sheet = "Assumptions"

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 22

    acq = pair["acquirer"]
    tgt = pair["target"]
    acq_v = pair_verdict["acquirer"]
    tgt_v = pair_verdict["target"]

    w = RowWriter(ws, sheet)
    refs = {}

    # ==========================================================================
    # SECTION A: COMPANY DATA
    # ==========================================================================
    w.section("SECTION A: COMPANY DATA (Fetched from SEC EDGAR)")

    # Column headers
    _put_value(ws, w.row, 3, "ACQUIRER", BLACK_FONT_BOLD)
    _put_value(ws, w.row, 4, "TARGET", BLACK_FONT_BOLD)
    w.row += 1

    def company_row(label, key, number_format=None, fmt=None):
        """Write one label row with acquirer + target fetched values."""
        _put_value(ws, w.row, 2, label, BLUE_FONT)
        for col_idx, profile in ((3, acq), (4, tgt)):
            value = profile.get(key)
            _put_value(ws, w.row, col_idx,
                       value if value is not None else "",
                       BLUE_FONT, number_format)
        w.row += 1

    company_row("Company Name", "company_name")
    company_row("Ticker", "ticker")

    _put_value(ws, w.row, 2, "SIC Code / Sector", BLUE_FONT)
    for col_idx, profile in ((3, acq), (4, tgt)):
        _put_value(ws, w.row, col_idx,
                   f"{profile.get('sic_code', '')} - {profile.get('sic_description', '')}",
                   BLUE_FONT)
    w.row += 1
    w.blank()

    # --- Market data ---
    w.subsection("Market Data")

    _put_value(ws, w.row, 2, "Current Share Price", BLUE_FONT)
    for col_idx, profile, role in ((3, acq, "acq"), (4, tgt, "tgt")):
        price = profile.get("current_price")
        if price is not None:
            _put_value(ws, w.row, col_idx, price, BLUE_FONT, PRICE_FORMAT)
        else:
            _put_value(ws, w.row, col_idx, 0, BLUE_FONT, PRICE_FORMAT,
                       fill=DEFAULTED_FILL,
                       comment="Price not available - update manually")
        refs[f"{role}_price"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
    _define_name(wb, "AcquirerPrice", sheet, f"C{w.row}")
    _define_name(wb, "TargetPrice", sheet, f"D{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "Shares Outstanding (basic)", BLUE_FONT)
    for col_idx, profile, role in ((3, acq, "acq"), (4, tgt, "tgt")):
        shares = profile.get("shares_outstanding")
        if shares is not None:
            _put_value(ws, w.row, col_idx, shares, BLUE_FONT, NUMBER_FORMAT)
        else:
            _put_value(ws, w.row, col_idx, 0, BLUE_FONT, NUMBER_FORMAT,
                       fill=DEFAULTED_FILL,
                       comment="Basic shares not available - update manually")
        refs[f"{role}_basic_shares"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
    w.row += 1

    _put_value(ws, w.row, 2, "Diluted Shares (weighted avg)", BLUE_FONT)
    for col_idx, profile, role in ((3, acq, "acq"), (4, tgt, "tgt")):
        _, fy_data = _most_recent_fy_data(profile)
        diluted = fy_data.get("diluted_shares") or profile.get("diluted_shares")
        if diluted is not None:
            _put_value(ws, w.row, col_idx, diluted, BLUE_FONT, NUMBER_FORMAT)
        else:
            basic = profile.get("shares_outstanding")
            if basic is not None:
                _put_value(ws, w.row, col_idx, basic, BLUE_FONT, NUMBER_FORMAT,
                           fill=DEFAULTED_FILL,
                           comment="Diluted share count unavailable - using "
                                   "basic shares outstanding as substitute")
            else:
                _put_value(ws, w.row, col_idx, 0, BLUE_FONT, NUMBER_FORMAT,
                           fill=DEFAULTED_FILL,
                           comment="No share count available - update manually")
        refs[f"{role}_diluted_shares"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
    _define_name(wb, "AcquirerDilutedShares", sheet, f"C{w.row}")
    _define_name(wb, "TargetDilutedShares", sheet, f"D{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "Diluted EPS", BLUE_FONT)
    for col_idx, profile, role in ((3, acq, "acq"), (4, tgt, "tgt")):
        _, fy_data = _most_recent_fy_data(profile)
        eps = fy_data.get("diluted_eps")
        if eps is not None:
            _put_value(ws, w.row, col_idx, eps, BLUE_FONT, EPS_FORMAT)
        else:
            _put_value(ws, w.row, col_idx, 0, BLUE_FONT, EPS_FORMAT,
                       fill=DEFAULTED_FILL,
                       comment="Diluted EPS unavailable - derived per-share "
                               "figures rely on Net Income / share count")
        refs[f"{role}_diluted_eps"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
    w.row += 1

    _put_value(ws, w.row, 2, "Market Cap", BLACK_FONT)
    for role, col_idx in (("acq", 3), ("tgt", 4)):
        col = chr(64 + col_idx)
        cell = _put_value(
            ws, w.row, col_idx,
            f"={refs[f'{role}_price']}*{refs[f'{role}_diluted_shares']}",
            BLACK_FONT, CURRENCY_FORMAT)
    refs["acq_market_cap"] = _abs_ref(sheet, 'C', w.row)
    refs["tgt_market_cap"] = _abs_ref(sheet, 'D', w.row)
    w.row += 1
    w.blank()

    # --- Most recent fiscal year financials ---
    w.subsection("Most Recent Fiscal Year Financials")

    acq_fy, acq_fy_data = _most_recent_fy_data(acq)
    tgt_fy, tgt_fy_data = _most_recent_fy_data(tgt)

    _put_value(ws, w.row, 2, "Fiscal Year", BLUE_FONT)
    _put_value(ws, w.row, 3, acq_fy or "", BLUE_FONT)
    _put_value(ws, w.row, 4, tgt_fy or "", BLUE_FONT)
    w.row += 1

    def financial_row(label, key, ref_key, defaulted_msgs=None):
        """One fetched financial row for both companies, yellow if defaulted."""
        _put_value(ws, w.row, 2, label, BLUE_FONT)
        for col_idx, fy_data, verdict, role in (
                (3, acq_fy_data, acq_v, "acq"), (4, tgt_fy_data, tgt_v, "tgt")):
            value = fy_data.get(key)
            if value is not None:
                _put_value(ws, w.row, col_idx, value, BLUE_FONT, CURRENCY_FORMAT)
            else:
                _put_value(ws, w.row, col_idx, 0, BLUE_FONT, CURRENCY_FORMAT,
                           fill=DEFAULTED_FILL,
                           comment=(defaulted_msgs or
                                    f"{label} unavailable - defaulted to 0, "
                                    f"verify before relying on downstream math"))
            refs[f"{role}_{ref_key}"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
        w.row += 1

    financial_row("Revenue", "revenue", "revenue")
    financial_row("Operating Income", "operating_income", "op_income")
    financial_row("D&A", "da", "da")

    # EBITDA (formula)
    _put_value(ws, w.row, 2, "EBITDA (Op Inc + D&A)", BLACK_FONT)
    for role, col_idx in (("acq", 3), ("tgt", 4)):
        _put_value(ws, w.row, col_idx,
                   f"={refs[f'{role}_op_income']}+{refs[f'{role}_da']}",
                   BLACK_FONT, CURRENCY_FORMAT)
        refs[f"{role}_ebitda"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
    w.row += 1

    financial_row("Net Income", "net_income", "net_income")
    financial_row("Total Debt", "total_debt", "total_debt",
                  "Total debt defaulted to $0 - verify this is accurate")
    financial_row("Cash & Equivalents", "cash", "cash",
                  "Cash defaulted to $0 - verify this is accurate")
    financial_row("Total Assets", "total_assets", "total_assets",
                  "Total assets unavailable - book equity for PPA will be "
                  "approximate (Goodwill precision degraded)")
    financial_row("Total Liabilities", "total_liabilities", "total_liabilities",
                  "Total liabilities unavailable - book equity for PPA will "
                  "be approximate (Goodwill precision degraded)")
    w.blank()

    # ==========================================================================
    # SECTION B: DEAL ASSUMPTIONS (user inputs, blue font)
    # ==========================================================================
    w.section("SECTION B: DEAL ASSUMPTIONS (User Inputs)")

    deal = dict(DEFAULT_DEAL_ASSUMPTIONS)
    if deal_assumptions:
        unknown = set(deal_assumptions) - set(deal)
        if unknown:
            raise ValueError(f"Unknown deal assumption keys: {sorted(unknown)}")
        deal.update(deal_assumptions)

    def deal_input(label, default_value, ref_key, number_format=PERCENT_FORMAT,
                   comment=None, name=None):
        _put_value(ws, w.row, 2, label, BLUE_FONT)
        _put_value(ws, w.row, 3, deal.get(ref_key, default_value), BLUE_FONT,
                   number_format, comment=comment)
        refs[ref_key] = _abs_ref(sheet, 'C', w.row)
        if name:
            _define_name(wb, name, sheet, f"C{w.row}")
        w.row += 1

    deal_input("Offer Premium %", 0.25, "offer_premium",
               comment="Default 25% premium over target's current share price",
               name="OfferPremium")
    deal_input("% Cash (of total consideration)", 0.40, "pct_cash",
               comment="Share of the purchase funded with balance-sheet cash",
               name="PctCash")
    deal_input("% Debt (of total consideration)", 0.30, "pct_debt",
               comment="Share of the purchase funded with new acquisition debt",
               name="PctDebt")

    # % Stock is NEVER an input — always the plug (100% - %Cash - %Debt)
    _put_value(ws, w.row, 2, "% Stock (plug = 100% - Cash - Debt)", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=1-{refs['pct_cash']}-{refs['pct_debt']}",
               BLACK_FONT, PERCENT_FORMAT)
    refs["pct_stock"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "PctStock", sheet, f"C{w.row}")
    w.row += 1

    deal_input("New Debt Interest Rate", 0.08, "new_debt_rate",
               comment="Interest rate on newly-raised acquisition debt",
               name="NewDebtRate")
    deal_input("Foregone Interest Rate on Cash Used", 0.045, "foregone_cash_rate",
               comment="Yield the acquirer stops earning on cash spent in the deal",
               name="ForegoneCashRate")
    deal_input("Revenue Synergies % (of combined revenue)", 0.0, "revenue_synergies_pct",
               comment="Default 0% - add deliberately, don't assume synergies exist",
               name="RevenueSynergiesPct")
    deal_input("Cost Synergies % (of combined opex)", 0.0, "cost_synergies_pct",
               comment="Default 0% - add deliberately, don't assume synergies exist",
               name="CostSynergiesPct")
    deal_input("Synergy Ramp Period (years to full run-rate)", 3,
               "synergy_ramp_years", number_format=YEARS_FORMAT,
               comment="Synergies phase in linearly over this many years",
               name="SynergyRampYears")
    deal_input("Transaction Fee %", 0.02, "transaction_fee_pct",
               comment="Advisory/financing fees as % of purchase equity value",
               name="TransactionFeePct")
    deal_input("Asset Step-Up % (of premium over book)", 0.50, "asset_step_up_pct",
               comment="ROUGH PLACEHOLDER - the true step-up depends on the "
                       "target's asset fair values; adjust based on real "
                       "diligence before relying on the PPA output",
               name="AssetStepUpPct")
    deal_input("Intangible Amortization Period (years)", 10,
               "intangible_amort_years", number_format=YEARS_FORMAT,
               comment="Stepped-up assets amortize straight-line over this period",
               name="IntangibleAmortYears")
    deal_input("Tax Rate", 0.25, "tax_rate",
               comment="Default 25%. SEC fetch does not include income tax "
                       "expense, so the acquirer's actual effective rate is "
                       "not derivable here - adjust to the acquirer's "
                       "effective rate from its 10-K if needed",
               name="TaxRate")
    deal_input("Analysis Years", 5, "analysis_years",
               number_format=YEARS_FORMAT,
               comment="Projection horizon for the pro forma model",
               name="AnalysisYears")
    w.blank()

    # ==========================================================================
    # SECTION C: PROJECTION ASSUMPTIONS (derived defaults, editable)
    # ==========================================================================
    w.section("SECTION C: PROJECTION ASSUMPTIONS (Derived Defaults)")

    _put_value(ws, w.row, 3, "ACQUIRER", BLACK_FONT_BOLD)
    _put_value(ws, w.row, 4, "TARGET", BLACK_FONT_BOLD)
    w.row += 1

    # Revenue growth defaults: AIO LBO's median-of-YoY + sanity-band method.
    _put_value(ws, w.row, 2, "Revenue Growth % (annual)", BLUE_FONT)
    for col_idx, profile, role in ((3, acq, "acq"), (4, tgt, "tgt")):
        rate, is_fallback, rate_comment = calculate_revenue_growth_rate(profile)
        _put_value(ws, w.row, col_idx, rate, BLUE_FONT, PERCENT_FORMAT,
                   fill=IMPLAUSIBLE_CALC_FILL if is_fallback else None,
                   comment=rate_comment)
        refs[f"{role}_growth"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
    _define_name(wb, "AcquirerGrowth", sheet, f"C{w.row}")
    _define_name(wb, "TargetGrowth", sheet, f"D{w.row}")
    w.row += 1

    # Operating margins (formulas off fetched data, zero-guarded)
    _put_value(ws, w.row, 2, "Operating Margin (most recent FY)", BLACK_FONT)
    for role, col_idx in (("acq", 3), ("tgt", 4)):
        rev_ref = refs[f"{role}_revenue"]
        op_ref = refs[f"{role}_op_income"]
        _put_value(ws, w.row, col_idx,
                   f"=IF({rev_ref}=0,0,{op_ref}/{rev_ref})",
                   BLACK_FONT, PERCENT_FORMAT)
        refs[f"{role}_op_margin"] = _abs_ref(sheet, chr(64 + col_idx), w.row)
    w.row += 1

    return refs


# =============================================================================
# TAB 2: SOURCES & USES
# =============================================================================

def build_sources_uses_tab(wb: Workbook, refs: dict) -> dict:
    """
    Build the Sources & Uses tab. All formula, black font, direct
    cross-sheet references to Assumptions. Returns updated refs dict with
    this tab's key cells.
    """
    ws = wb.create_sheet("Sources & Uses", 1)
    sheet = "Sources & Uses"

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 22

    w = RowWriter(ws, sheet)

    # ==========================================================================
    # USES OF FUNDS
    # ==========================================================================
    w.section("USES OF FUNDS", span_end_col=3)

    _put_value(ws, w.row, 2, "Purchase Equity Value", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"={refs['tgt_price']}*(1+{refs['offer_premium']})*{refs['tgt_diluted_shares']}",
               BLACK_FONT, CURRENCY_FORMAT)
    pev_row = w.row
    refs["purchase_equity_value"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "PurchaseEquityValue", sheet, f"C{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "Transaction Fees", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{pev_row}*{refs['transaction_fee_pct']}",
               BLACK_FONT, CURRENCY_FORMAT)
    fees_row = w.row
    refs["transaction_fees"] = _abs_ref(sheet, 'C', w.row)
    w.row += 2

    _put_value(ws, w.row, 2, "Total Uses", BLACK_FONT_BOLD)
    total_uses_cell = _put_value(ws, w.row, 3,
                                 f"=C{pev_row}+C{fees_row}",
                                 BLACK_FONT_BOLD, CURRENCY_FORMAT)
    total_uses_cell.border = Border(top=Side(style='thin'),
                                    bottom=Side(style='double'))
    total_uses_row = w.row
    refs["total_uses"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "TotalUses", sheet, f"C{w.row}")
    w.row += 3

    # ==========================================================================
    # SOURCES OF FUNDS
    # ==========================================================================
    w.section("SOURCES OF FUNDS", span_end_col=3)

    _put_value(ws, w.row, 2, "Cash Used", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{total_uses_row}*{refs['pct_cash']}",
               BLACK_FONT, CURRENCY_FORMAT)
    cash_row = w.row
    refs["cash_used"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "CashUsed", sheet, f"C{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "New Debt Raised", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{total_uses_row}*{refs['pct_debt']}",
               BLACK_FONT, CURRENCY_FORMAT)
    debt_row = w.row
    refs["new_debt"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "NewDebtRaised", sheet, f"C{w.row}")
    w.row += 1

    # New Stock Issued is the PLUG: Total Uses - Cash - Debt
    _put_value(ws, w.row, 2, "New Stock Issued ($) [plug]", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{total_uses_row}-C{cash_row}-C{debt_row}",
               BLACK_FONT, CURRENCY_FORMAT)
    stock_row = w.row
    refs["new_stock_issued"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "NewStockIssued", sheet, f"C{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "New Shares Issued", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{stock_row}/{refs['acq_price']}",
               BLACK_FONT, NUMBER_FORMAT)
    refs["new_shares_issued"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "NewSharesIssued", sheet, f"C{w.row}")
    w.row += 2

    _put_value(ws, w.row, 2, "Total Sources", BLACK_FONT_BOLD)
    total_sources_cell = _put_value(
        ws, w.row, 3,
        f"=C{cash_row}+C{debt_row}+C{stock_row}",
        BLACK_FONT_BOLD, CURRENCY_FORMAT)
    total_sources_cell.border = Border(top=Side(style='thin'),
                                       bottom=Side(style='double'))
    total_sources_row = w.row
    refs["total_sources"] = _abs_ref(sheet, 'C', w.row)
    w.row += 3

    # ==========================================================================
    # BALANCE CHECK
    # ==========================================================================
    w.subsection("Balance Check", span_end_col=3)

    _put_value(ws, w.row, 2, "Sources - Uses (should = 0)", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{total_sources_row}-C{total_uses_row}",
               BLACK_FONT, CURRENCY_FORMAT)
    balance_row = w.row
    refs["balance_check"] = _abs_ref(sheet, 'C', w.row)

    ws.conditional_formatting.add(
        f"C{balance_row}",
        FormulaRule(formula=[f"C{balance_row}<>0"], fill=ERROR_FILL)
    )
    _define_name(wb, "BalanceCheck", sheet, f"C{balance_row}")

    return refs


# =============================================================================
# TAB 3: PURCHASE PRICE ALLOCATION
# =============================================================================

def build_ppa_tab(wb: Workbook, refs: dict, pair_verdict: dict) -> dict:
    """
    Build the Purchase Price Allocation tab. All formula, black font,
    direct cross-sheet references only.
    """
    ws = wb.create_sheet("Purchase Price Allocation", 2)
    sheet = "Purchase Price Allocation"

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 22

    w = RowWriter(ws, sheet)
    goodwill_degraded = pair_verdict["target"].get("goodwill_precision_degraded", False)

    w.section("PURCHASE PRICE ALLOCATION", span_end_col=3)

    _put_value(ws, w.row, 2, "Target Book Equity (Assets - Liabilities)", BLACK_FONT)
    book_equity_comment = None
    if goodwill_degraded:
        book_equity_comment = (
            "Target Total Assets/Liabilities were unavailable from SEC data - "
            "book equity (and therefore Goodwill) is APPROXIMATE. Enter real "
            "balance sheet values on the Assumptions tab to fix.")
    _put_value(ws, w.row, 3,
               f"={refs['tgt_total_assets']}-{refs['tgt_total_liabilities']}",
               BLACK_FONT, CURRENCY_FORMAT, comment=book_equity_comment)
    book_equity_row = w.row
    refs["tgt_book_equity"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "TargetBookEquity", sheet, f"C{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "Premium Over Book", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"={refs['purchase_equity_value']}-C{book_equity_row}",
               BLACK_FONT, CURRENCY_FORMAT)
    premium_row = w.row
    refs["premium_over_book"] = _abs_ref(sheet, 'C', w.row)
    w.row += 1

    _put_value(ws, w.row, 2, "Asset Step-Up", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{premium_row}*{refs['asset_step_up_pct']}",
               BLACK_FONT, CURRENCY_FORMAT)
    stepup_row = w.row
    refs["asset_step_up"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "AssetStepUp", sheet, f"C{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "Goodwill", BLACK_FONT)
    goodwill_comment = None
    if goodwill_degraded:
        goodwill_comment = ("APPROXIMATE - target book equity unavailable "
                            "(see Assumptions tab)")
    _put_value(ws, w.row, 3,
               f"=C{premium_row}-C{stepup_row}",
               BLACK_FONT, CURRENCY_FORMAT, comment=goodwill_comment)
    refs["goodwill"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "Goodwill", sheet, f"C{w.row}")
    w.row += 2

    _put_value(ws, w.row, 2, "Incremental Annual D&A (from step-up)", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=C{stepup_row}/{refs['intangible_amort_years']}",
               BLACK_FONT, CURRENCY_FORMAT)
    refs["incremental_da"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "IncrementalDA", sheet, f"C{w.row}")

    return refs


# =============================================================================
# TAB 4: PRO FORMA INCOME STATEMENT
# =============================================================================

def build_proforma_tab(wb: Workbook, refs: dict, analysis_years: int) -> dict:
    """
    Build the Pro Forma Income Statement tab. Year columns 1..analysis_years
    (dynamically sized off the Analysis Years input, not hardcoded). All
    formula, black font, direct cross-sheet references only.

    Combined Diluted Shares is deliberately FIXED across years: only the
    stock portion of consideration adds shares, at close. No buybacks.
    """
    ws = wb.create_sheet("Pro Forma IS", 3)
    sheet = "Pro Forma IS"

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 40
    for i in range(analysis_years):
        ws.column_dimensions[chr(67 + i)].width = 18

    w = RowWriter(ws, sheet)
    w.section("PRO FORMA INCOME STATEMENT (Combined Entity)",
              span_end_col=2 + analysis_years)

    def col(t):
        """Column letter for year t (1-based)."""
        return chr(66 + t)  # year 1 -> C

    # Year header row
    _put_value(ws, w.row, 2, "Year", BLACK_FONT_BOLD)
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t, t, BLACK_FONT_BOLD, YEARS_FORMAT)
    year_hdr_row = w.row
    w.row += 2

    # --- Revenue build ---
    w.subsection("Revenue", span_end_col=2 + analysis_years)

    _put_value(ws, w.row, 2, "Acquirer Revenue", BLACK_FONT)
    acq_rev_row = w.row
    for t in range(1, analysis_years + 1):
        if t == 1:
            formula = f"={refs['acq_revenue']}*(1+{refs['acq_growth']})"
        else:
            formula = f"={col(t-1)}{acq_rev_row}*(1+{refs['acq_growth']})"
        _put_value(ws, w.row, 2 + t, formula, BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Target Revenue", BLACK_FONT)
    tgt_rev_row = w.row
    for t in range(1, analysis_years + 1):
        if t == 1:
            formula = f"={refs['tgt_revenue']}*(1+{refs['tgt_growth']})"
        else:
            formula = f"={col(t-1)}{tgt_rev_row}*(1+{refs['tgt_growth']})"
        _put_value(ws, w.row, 2 + t, formula, BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Combined Revenue", BLACK_FONT_BOLD)
    comb_rev_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={col(t)}{acq_rev_row}+{col(t)}{tgt_rev_row}",
                   BLACK_FONT_BOLD, CURRENCY_FORMAT)
    w.row += 2

    # --- EBIT build ---
    w.subsection("EBIT", span_end_col=2 + analysis_years)

    _put_value(ws, w.row, 2, "Acquirer EBIT (margin held constant)", BLACK_FONT)
    acq_ebit_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={col(t)}{acq_rev_row}*{refs['acq_op_margin']}",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Target EBIT (margin held constant)", BLACK_FONT)
    tgt_ebit_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={col(t)}{tgt_rev_row}*{refs['tgt_op_margin']}",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    # Combined opex (needed as the cost-synergy base): Revenue - EBIT
    _put_value(ws, w.row, 2, "Combined Operating Expenses", BLACK_FONT)
    comb_opex_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={col(t)}{comb_rev_row}-{col(t)}{acq_ebit_row}-{col(t)}{tgt_ebit_row}",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    # Synergies: Full Run-Rate x MIN(1, t / Ramp Period) — linear ramp
    _put_value(ws, w.row, 2, "Revenue Synergies Realized", BLACK_FONT)
    rev_syn_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(
            ws, w.row, 2 + t,
            f"={col(t)}{comb_rev_row}*{refs['revenue_synergies_pct']}"
            f"*MIN(1,{col(t)}{year_hdr_row}/{refs['synergy_ramp_years']})",
            BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Cost Synergies Realized", BLACK_FONT)
    cost_syn_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(
            ws, w.row, 2 + t,
            f"={col(t)}{comb_opex_row}*{refs['cost_synergies_pct']}"
            f"*MIN(1,{col(t)}{year_hdr_row}/{refs['synergy_ramp_years']})",
            BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Incremental D&A (from PPA)", BLACK_FONT)
    inc_da_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t, f"={refs['incremental_da']}",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Combined EBIT", BLACK_FONT_BOLD)
    comb_ebit_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(
            ws, w.row, 2 + t,
            f"={col(t)}{acq_ebit_row}+{col(t)}{tgt_ebit_row}"
            f"+{col(t)}{rev_syn_row}+{col(t)}{cost_syn_row}-{col(t)}{inc_da_row}",
            BLACK_FONT_BOLD, CURRENCY_FORMAT)
    w.row += 2

    # --- Financing effects and taxes ---
    w.subsection("Financing & Taxes", span_end_col=2 + analysis_years)

    _put_value(ws, w.row, 2, "New Interest Expense", BLACK_FONT)
    interest_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={refs['new_debt']}*{refs['new_debt_rate']}",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Foregone Interest Income", BLACK_FONT)
    foregone_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={refs['cash_used']}*{refs['foregone_cash_rate']}",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Combined Pre-Tax Income", BLACK_FONT)
    pretax_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={col(t)}{comb_ebit_row}-{col(t)}{interest_row}-{col(t)}{foregone_row}",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    # Taxes floored at 0 — a loss-making combined entity is not a refund
    _put_value(ws, w.row, 2, "Taxes (floored at 0)", BLACK_FONT)
    taxes_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"=MAX(0,{col(t)}{pretax_row}*{refs['tax_rate']})",
                   BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Combined Net Income", BLACK_FONT_BOLD)
    comb_ni_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={col(t)}{pretax_row}-{col(t)}{taxes_row}",
                   BLACK_FONT_BOLD, CURRENCY_FORMAT)
    w.row += 2

    # --- Per share ---
    w.subsection("Per Share", span_end_col=2 + analysis_years)

    # FIXED across all years: acquirer diluted shares + new shares at close
    _put_value(ws, w.row, 2, "Combined Diluted Shares (fixed)", BLACK_FONT)
    comb_shares_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t,
                   f"={refs['acq_diluted_shares']}+{refs['new_shares_issued']}",
                   BLACK_FONT, NUMBER_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Pro Forma EPS", BLACK_FONT_BOLD)
    proforma_eps_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(
            ws, w.row, 2 + t,
            f"=IF({col(t)}{comb_shares_row}=0,0,"
            f"{col(t)}{comb_ni_row}/{col(t)}{comb_shares_row})",
            BLACK_FONT_BOLD, EPS_FORMAT)
    w.row += 1

    refs["proforma_rows"] = {
        "year_hdr_row": year_hdr_row,
        "acq_rev_row": acq_rev_row,
        "tgt_rev_row": tgt_rev_row,
        "comb_rev_row": comb_rev_row,
        "acq_ebit_row": acq_ebit_row,
        "tgt_ebit_row": tgt_ebit_row,
        "comb_opex_row": comb_opex_row,
        "rev_syn_row": rev_syn_row,
        "cost_syn_row": cost_syn_row,
        "inc_da_row": inc_da_row,
        "comb_ebit_row": comb_ebit_row,
        "interest_row": interest_row,
        "foregone_row": foregone_row,
        "pretax_row": pretax_row,
        "taxes_row": taxes_row,
        "comb_ni_row": comb_ni_row,
        "comb_shares_row": comb_shares_row,
        "proforma_eps_row": proforma_eps_row,
    }
    # Direct refs to the per-year Pro Forma EPS cells for the A/D tab
    refs["proforma_eps_cells"] = [
        _abs_ref(sheet, col(t), proforma_eps_row)
        for t in range(1, analysis_years + 1)
    ]
    refs["comb_shares_cells"] = [
        _abs_ref(sheet, col(t), comb_shares_row)
        for t in range(1, analysis_years + 1)
    ]
    return refs


# =============================================================================
# TAB 5: ACCRETION / DILUTION
# =============================================================================

def build_accretion_tab(wb: Workbook, refs: dict, analysis_years: int) -> dict:
    """
    Build the Accretion/Dilution tab: headline summary at top (Year 1 and
    final-year A/D % side by side, large font), then the full year-by-year
    trajectory. Direct cross-sheet references only.
    """
    ws = wb.create_sheet("Accretion Dilution", 4)
    sheet = "Accretion Dilution"

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 40
    for i in range(max(analysis_years, 2)):
        ws.column_dimensions[chr(67 + i)].width = 18

    w = RowWriter(ws, sheet)

    def col(t):
        return chr(66 + t)

    # ==========================================================================
    # RESULTS SUMMARY — the headline numbers, top of tab, large font
    # ==========================================================================
    w.section("RESULTS SUMMARY", span_end_col=2 + max(analysis_years, 2))

    _put_value(ws, w.row, 3, "Year 1", HEADLINE_LABEL_FONT)
    _put_value(ws, w.row, 4, f"Year {analysis_years}", HEADLINE_LABEL_FONT)
    w.row += 1

    _put_value(ws, w.row, 2, "Accretion / (Dilution) %", HEADLINE_LABEL_FONT)
    headline_row = w.row
    # These formulas reference the A/D trajectory row built below (same
    # sheet) — the row number is deterministic, computed after layout below.
    w.row += 2

    # ==========================================================================
    # STANDALONE BASELINE
    # ==========================================================================
    w.subsection("Acquirer Standalone (actual reported figures)",
                 span_end_col=2 + max(analysis_years, 2))

    _put_value(ws, w.row, 2, "Acquirer Net Income (fetched)", BLACK_FONT)
    _put_value(ws, w.row, 3, f"={refs['acq_net_income']}",
               BLACK_FONT, CURRENCY_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Acquirer Diluted Shares (fetched)", BLACK_FONT)
    _put_value(ws, w.row, 3, f"={refs['acq_diluted_shares']}",
               BLACK_FONT, NUMBER_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Acquirer Standalone EPS", BLACK_FONT_BOLD)
    _put_value(
        ws, w.row, 3,
        f"=IF({refs['acq_diluted_shares']}=0,0,"
        f"{refs['acq_net_income']}/{refs['acq_diluted_shares']})",
        BLACK_FONT_BOLD, EPS_FORMAT)
    standalone_eps_row = w.row
    refs["standalone_eps"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "StandaloneEPS", sheet, f"C{w.row}")
    w.row += 2

    # ==========================================================================
    # YEAR-BY-YEAR TRAJECTORY — the visual focus of this tab
    # ==========================================================================
    w.subsection("Accretion / (Dilution) Trajectory",
                 span_end_col=2 + max(analysis_years, 2))

    _put_value(ws, w.row, 2, "Year", BLACK_FONT_BOLD)
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t, t, BLACK_FONT_BOLD, YEARS_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Pro Forma EPS", BLACK_FONT)
    proforma_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t, f"={refs['proforma_eps_cells'][t-1]}",
                   BLACK_FONT, EPS_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Acquirer Standalone EPS", BLACK_FONT)
    standalone_traj_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(ws, w.row, 2 + t, f"=$C${standalone_eps_row}",
                   BLACK_FONT, EPS_FORMAT)
    w.row += 1

    _put_value(ws, w.row, 2, "Accretion / (Dilution) %", BLACK_FONT_BOLD)
    ad_row = w.row
    for t in range(1, analysis_years + 1):
        _put_value(
            ws, w.row, 2 + t,
            f"=IF({col(t)}{standalone_traj_row}=0,0,"
            f"{col(t)}{proforma_row}/{col(t)}{standalone_traj_row}-1)",
            BLACK_FONT_BOLD, PERCENT_FORMAT)
        # Red fill when dilutive
        ws.conditional_formatting.add(
            f"{col(t)}{ad_row}",
            FormulaRule(formula=[f"{col(t)}{ad_row}<0"], fill=ERROR_FILL)
        )
    w.row += 1

    refs["ad_cells"] = [
        _abs_ref(sheet, col(t), ad_row) for t in range(1, analysis_years + 1)
    ]
    refs["ad_year1"] = refs["ad_cells"][0]
    refs["ad_final"] = refs["ad_cells"][-1]

    # Now fill in the headline summary cells (same sheet, so plain refs)
    _put_value(ws, headline_row, 3, f"={col(1)}{ad_row}",
               HEADLINE_FONT, PERCENT_FORMAT)
    _put_value(ws, headline_row, 4, f"={col(analysis_years)}{ad_row}",
               HEADLINE_FONT, PERCENT_FORMAT)
    for hc in ('C', 'D'):
        ws.conditional_formatting.add(
            f"{hc}{headline_row}",
            FormulaRule(formula=[f"{hc}{headline_row}<0"], fill=ERROR_FILL)
        )
    refs["headline_year1"] = _abs_ref(sheet, 'C', headline_row)
    refs["headline_final"] = _abs_ref(sheet, 'D', headline_row)

    return refs


# =============================================================================
# TAB 6: MARKET CONCENTRATION (HHI)
# =============================================================================

# Number of blank competitor rows below the pre-populated acquirer/target
COMPETITOR_ROWS = 8

def build_concentration_tab(wb: Workbook, refs: dict, pair: dict) -> dict:
    """
    Build the Market Concentration tab. Deliberately split between
    deterministic auto-filled content and clearly-labeled human/AI-assisted
    judgment (user-entered market shares tagged user_provided).

    HHI uses percentage-point convention (40% share contributes 40^2=1600),
    with shares entered as percent-formatted fractions and formulas
    multiplying by 10,000.
    """
    ws = wb.create_sheet("Market Concentration", 5)
    sheet = "Market Concentration"

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['C'].width = 26
    ws.column_dimensions['D'].width = 26

    acq = pair["acquirer"]
    tgt = pair["target"]

    w = RowWriter(ws, sheet)

    # ==========================================================================
    # AUTO-FILLED, DETERMINISTIC
    # ==========================================================================
    w.section("INDUSTRY CLASSIFICATION (Auto-filled from SEC EDGAR)")

    _put_value(ws, w.row, 3, "ACQUIRER", BLACK_FONT_BOLD)
    _put_value(ws, w.row, 4, "TARGET", BLACK_FONT_BOLD)
    w.row += 1

    _put_value(ws, w.row, 2, "Company", BLUE_FONT)
    _put_value(ws, w.row, 3, acq.get("company_name", ""), BLUE_FONT)
    _put_value(ws, w.row, 4, tgt.get("company_name", ""), BLUE_FONT)
    w.row += 1

    _put_value(ws, w.row, 2, "SIC Code", BLUE_FONT)
    _put_value(ws, w.row, 3, str(acq.get("sic_code", "")), BLUE_FONT)
    _put_value(ws, w.row, 4, str(tgt.get("sic_code", "")), BLUE_FONT)
    sic_row = w.row
    w.row += 1

    _put_value(ws, w.row, 2, "Industry Description", BLUE_FONT)
    _put_value(ws, w.row, 3, acq.get("sic_description", ""), BLUE_FONT)
    _put_value(ws, w.row, 4, tgt.get("sic_description", ""), BLUE_FONT)
    w.row += 1

    _put_value(ws, w.row, 2, "Same Industry (same SIC code)", BLACK_FONT)
    _put_value(ws, w.row, 3, f"=IF(C{sic_row}=D{sic_row},TRUE,FALSE)",
               BLACK_FONT)
    refs["same_industry_flag"] = _abs_ref(sheet, 'C', w.row)
    w.row += 2

    # ==========================================================================
    # RESEARCH-ASSIST PROMPT (copyable, run externally by the USER — this
    # tool never makes the research call itself)
    # ==========================================================================
    w.subsection("Research Assist — copy this prompt to any LLM with web search")

    prompt_text = (
        f"Research current market share estimates for a potential merger "
        f"analysis. The two companies are: (1) {acq.get('company_name', '?')} "
        f"(ticker {acq.get('ticker', '?')}) and (2) {tgt.get('company_name', '?')} "
        f"(ticker {tgt.get('ticker', '?')}). Both operate in or around the "
        f"industry: \"{tgt.get('sic_description', 'Unknown')}\" (SEC SIC "
        f"{tgt.get('sic_code', '?')}; acquirer SIC {acq.get('sic_code', '?')} - "
        f"\"{acq.get('sic_description', 'Unknown')}\"). Please: (a) define the "
        f"most sensible relevant market these two actually compete in, "
        f"(b) list the major competitors in that market, and (c) give "
        f"best-available estimated market share percentages for each company "
        f"including the two above, with sources and dates for every estimate. "
        f"Flag any number that is an inference rather than a published figure."
    )
    prompt_cell = _put_value(ws, w.row, 2, prompt_text, BLACK_FONT,
                             comment="Copy this text into an external LLM "
                                     "with search. Review its answers, then "
                                     "type the market shares you trust into "
                                     "the table below. AI research stays "
                                     "human-reviewed - it is never a silent "
                                     "source of truth in this model.")
    prompt_cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=w.row, start_column=2, end_row=w.row + 3,
                   end_column=4)
    refs["research_prompt"] = _abs_ref(sheet, 'B', w.row)
    w.row += 5

    # ==========================================================================
    # USER-INPUT MARKET SHARE TABLE (provenance tag: user_provided)
    # ==========================================================================
    w.subsection("Market Share Estimates (USER INPUT - light green area)")

    _put_value(ws, w.row, 2, "Company Name", BLACK_FONT_BOLD)
    _put_value(ws, w.row, 3, "Estimated Market Share %", BLACK_FONT_BOLD)
    w.row += 1

    table_start = w.row
    name_cells, share_cells = [], []

    def share_input_row(prefill_name=None, note=None):
        name_cell = ws.cell(row=w.row, column=2)
        share_cell = ws.cell(row=w.row, column=3)
        if prefill_name:
            name_cell.value = prefill_name
        name_cell.font = BLUE_FONT
        share_cell.font = BLUE_FONT
        share_cell.number_format = PERCENT_FORMAT
        name_cell.fill = USER_PROVIDED_FILL
        share_cell.fill = USER_PROVIDED_FILL
        if note:
            share_cell.comment = Comment(note, "System")
        name_cells.append(f"B{w.row}")
        share_cells.append(f"C{w.row}")
        w.row += 1

    share_input_row(acq.get("company_name", "Acquirer"),
                    "user_provided: enter this company's estimated market "
                    "share (e.g. 25%)")
    acq_share_cell = share_cells[0]
    share_input_row(tgt.get("company_name", "Target"),
                    "user_provided: enter this company's estimated market "
                    "share (e.g. 10%)")
    tgt_share_cell = share_cells[1]
    for _ in range(COMPETITOR_ROWS):
        share_input_row()
    table_end = w.row - 1

    all_share_range = f"C{table_start}:C{table_end}"
    comp_share_range = f"C{table_start + 2}:C{table_end}"  # competitors only

    refs["market_share_table"] = {
        "sheet": sheet,
        "name_cells": name_cells,
        "share_cells": share_cells,
        "acq_share_cell": _abs_ref(sheet, 'C', table_start),
        "tgt_share_cell": _abs_ref(sheet, 'C', table_start + 1),
        "provenance": "user_provided",
    }
    w.blank()

    # ==========================================================================
    # DETERMINISTIC HHI CALCULATION
    # ==========================================================================
    w.subsection("HHI Calculation (deterministic once table is filled)")

    EMPTY_MSG = "not yet calculated — fill in market share estimates above"
    has_data = f"COUNT({all_share_range})>0"

    _put_value(ws, w.row, 2, "Pre-Merger HHI (all companies separate)", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=IF({has_data},SUMPRODUCT({all_share_range},{all_share_range})*10000,"
               f"\"{EMPTY_MSG}\")",
               BLACK_FONT, NUMBER_FORMAT)
    pre_hhi_row = w.row
    refs["pre_merger_hhi"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "PreMergerHHI", sheet, f"C{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "Post-Merger HHI (acquirer+target combined)", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=IF({has_data},"
               f"(({acq_share_cell}+{tgt_share_cell})^2)*10000"
               f"+SUMPRODUCT({comp_share_range},{comp_share_range})*10000,"
               f"\"{EMPTY_MSG}\")",
               BLACK_FONT, NUMBER_FORMAT)
    post_hhi_row = w.row
    refs["post_merger_hhi"] = _abs_ref(sheet, 'C', w.row)
    _define_name(wb, "PostMergerHHI", sheet, f"C{w.row}")
    w.row += 1

    _put_value(ws, w.row, 2, "Delta HHI (post - pre)", BLACK_FONT)
    _put_value(ws, w.row, 3,
               f"=IF({has_data},C{post_hhi_row}-C{pre_hhi_row},\"{EMPTY_MSG}\")",
               BLACK_FONT, NUMBER_FORMAT)
    delta_hhi_row = w.row
    refs["delta_hhi"] = _abs_ref(sheet, 'C', w.row)
    w.row += 2

    # DOJ/FTC threshold classifications — formula-driven, not static text
    def classification_formula(hhi_cell):
        return (f"=IF(NOT(ISNUMBER({hhi_cell})),\"—\","
                f"IF({hhi_cell}<1500,\"Unconcentrated\","
                f"IF({hhi_cell}<=2500,\"Moderately Concentrated\","
                f"\"Highly Concentrated\")))")

    _put_value(ws, w.row, 2, "Pre-Merger Classification (DOJ/FTC)", BLACK_FONT)
    _put_value(ws, w.row, 3, classification_formula(f"C{pre_hhi_row}"),
               BLACK_FONT)
    refs["pre_merger_class"] = _abs_ref(sheet, 'C', w.row)
    w.row += 1

    _put_value(ws, w.row, 2, "Post-Merger Classification (DOJ/FTC)", BLACK_FONT)
    _put_value(ws, w.row, 3, classification_formula(f"C{post_hhi_row}"),
               BLACK_FONT)
    refs["post_merger_class"] = _abs_ref(sheet, 'C', w.row)
    w.row += 1

    # The REAL presumptive-concern trigger: Delta > 200 AND Post > 2,500
    _put_value(ws, w.row, 2,
               "Presumptive Concern (Delta > 200 AND Post-Merger > 2,500)",
               BLACK_FONT)
    _put_value(
        ws, w.row, 3,
        f"=IF(NOT(ISNUMBER(C{post_hhi_row})),\"—\","
        f"IF(AND(C{delta_hhi_row}>200,C{post_hhi_row}>2500),"
        f"\"PRESUMED LIKELY TO ENHANCE MARKET POWER\","
        f"\"Below presumptive-concern trigger\"))",
        BLACK_FONT)
    presumption_row = w.row
    refs["presumptive_concern"] = _abs_ref(sheet, 'C', w.row)
    ws.conditional_formatting.add(
        f"C{presumption_row}",
        FormulaRule(formula=[f'C{presumption_row}="PRESUMED LIKELY TO ENHANCE MARKET POWER"'],
                    fill=ERROR_FILL)
    )

    return refs


# =============================================================================
# TAB 7: SENSITIVITY
# =============================================================================

def build_sensitivity_tab(wb: Workbook, refs: dict, deal: dict) -> dict:
    """
    Build the Sensitivity tab: Year-1 Accretion/Dilution % across
    Offer Premium (rows) x Financing Mix (columns).

    SECOND AXIS DEFINITION (documented per the build prompt): the column
    axis is % STOCK of total consideration. The non-stock remainder is
    split between cash and debt in the BASE CASE's cash:debt ratio
    (e.g. base 40% cash / 30% debt -> non-stock portion splits 4:3).

    NO Excel Data Table / What-If feature (unreliable in Google Sheets /
    LibreOffice — AIO LBO lesson). Every grid cell is a fully independent,
    self-contained formula that recomputes Purchase Equity Value through
    Pro Forma EPS for its own premium/mix combination, referencing fixed
    base-case values (Year-1 EBIT components, synergies) only where they
    genuinely don't vary with premium or financing mix.
    """
    ws = wb.create_sheet("Sensitivity", 6)
    sheet = "Sensitivity"

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 26

    w = RowWriter(ws, sheet)

    base_premium = deal["offer_premium"]
    base_stock = round(1.0 - deal["pct_cash"] - deal["pct_debt"], 10)

    # Axes include the exact base-case values so the base cell must
    # reproduce the Accretion/Dilution tab's Year-1 number exactly.
    premiums = sorted(set([0.10, 0.175, 0.325, 0.40, base_premium]))
    stock_mixes = sorted(set([0.0, 0.25, 0.50, 0.75, 1.0, base_stock]))

    for i in range(len(stock_mixes)):
        ws.column_dimensions[chr(67 + i)].width = 16

    w.section("SENSITIVITY: YEAR-1 ACCRETION/(DILUTION) %",
              span_end_col=2 + len(stock_mixes))

    _put_value(
        ws, w.row, 2,
        "Rows: Offer Premium %. Columns: % Stock of consideration; the "
        "non-stock remainder splits cash:debt in the base-case ratio "
        f"({deal['pct_cash']:.0%}:{deal['pct_debt']:.0%}). Each cell "
        "independently recomputes purchase price through pro forma EPS.",
        BLACK_FONT)
    ws.merge_cells(start_row=w.row, start_column=2, end_row=w.row,
                   end_column=2 + len(stock_mixes))
    ws.cell(row=w.row, column=2).alignment = Alignment(wrap_text=True)
    w.row += 2

    # Column headers (% stock)
    hdr_row = w.row
    _put_value(ws, w.row, 2, "Premium \\ % Stock", BLACK_FONT_BOLD)
    for i, s in enumerate(stock_mixes):
        cell = _put_value(ws, w.row, 3 + i, s, BLUE_FONT_BOLD, PERCENT_FORMAT)
        if abs(s - base_stock) < 1e-9:
            cell.comment = Comment("Base case financing mix", "System")
    w.row += 1

    # Fixed base-case Year-1 values that do NOT vary with premium/mix:
    pf = refs["proforma_rows"]
    y1 = 'C'  # year-1 column on the Pro Forma IS tab
    ebit_pre_da = (f"('Pro Forma IS'!${y1}${pf['acq_ebit_row']}"
                   f"+'Pro Forma IS'!${y1}${pf['tgt_ebit_row']}"
                   f"+'Pro Forma IS'!${y1}${pf['rev_syn_row']}"
                   f"+'Pro Forma IS'!${y1}${pf['cost_syn_row']})")

    grid_cells = {}
    grid_start_row = w.row

    for r, p in enumerate(premiums):
        prem_cell = _put_value(ws, w.row, 2, p, BLUE_FONT_BOLD, PERCENT_FORMAT)
        if abs(p - base_premium) < 1e-9:
            prem_cell.comment = Comment("Base case offer premium", "System")
        P = f"$B{w.row}"

        for c, s in enumerate(stock_mixes):
            col_letter = chr(67 + c)
            S = f"{col_letter}${hdr_row}"

            # Self-contained recomputation for this (premium, mix) cell:
            pev = f"({refs['tgt_price']}*(1+{P})*{refs['tgt_diluted_shares']})"
            uses = f"({pev}*(1+{refs['transaction_fee_pct']}))"
            stock_d = f"({uses}*{S})"
            new_sh = f"({stock_d}/{refs['acq_price']})"
            nonstock = f"({uses}*(1-{S}))"
            cash_ratio = (f"IF({refs['pct_cash']}+{refs['pct_debt']}=0,0.5,"
                          f"{refs['pct_cash']}/({refs['pct_cash']}+{refs['pct_debt']}))")
            cash_d = f"({nonstock}*{cash_ratio})"
            debt_d = f"({nonstock}*(1-{cash_ratio}))"
            inc_da = (f"((({pev}-({refs['tgt_total_assets']}-{refs['tgt_total_liabilities']}))"
                      f"*{refs['asset_step_up_pct']})/{refs['intangible_amort_years']})")
            pretax = (f"({ebit_pre_da}-{inc_da}"
                      f"-({debt_d}*{refs['new_debt_rate']})"
                      f"-({cash_d}*{refs['foregone_cash_rate']}))")
            ni = f"({pretax}-MAX(0,{pretax}*{refs['tax_rate']}))"
            shares = f"({refs['acq_diluted_shares']}+{new_sh})"
            eps = f"IF({shares}=0,0,{ni}/{shares})"
            formula = (f"=IF({refs['standalone_eps']}=0,0,"
                       f"{eps}/{refs['standalone_eps']}-1)")

            _put_value(ws, w.row, 3 + c, formula, BLACK_FONT, PERCENT_FORMAT)
            ws.conditional_formatting.add(
                f"{col_letter}{w.row}",
                FormulaRule(formula=[f"{col_letter}{w.row}<0"], fill=ERROR_FILL)
            )
            grid_cells[(p, s)] = _abs_ref(sheet, col_letter, w.row)
        w.row += 1

    refs["sensitivity_grid"] = {
        "premiums": premiums,
        "stock_mixes": stock_mixes,
        "base_premium": base_premium,
        "base_stock": base_stock,
        "cells": grid_cells,
        "base_cell": grid_cells[(base_premium, base_stock)],
    }
    return refs


# =============================================================================
# WORKBOOK ASSEMBLY
# =============================================================================

def generate_ma_workbook(pair: dict, pair_verdict: dict, output_path: str,
                         deal_assumptions: Optional[dict] = None) -> dict:
    """
    Generate the M&A workbook for a validated pair.

    Args:
        pair: result of fetch_ma_pair()
        pair_verdict: result of validate_pair() — caller should have
            already rejected status == "fail" pairs
        output_path: where to save the .xlsx
        deal_assumptions: optional overrides for DEFAULT_DEAL_ASSUMPTIONS
            (e.g. {"pct_cash": 0, "pct_debt": 0} for an all-stock deal)

    Returns the refs dict (semantic name -> direct cell reference) for
    downstream use/testing.
    """
    if pair_verdict.get("status") == "fail":
        raise ValueError(
            "Cannot generate a workbook for a failed validation: "
            + "; ".join(pair_verdict.get("disqualifying_reasons", []))
        )

    deal = dict(DEFAULT_DEAL_ASSUMPTIONS)
    if deal_assumptions:
        deal.update(deal_assumptions)
    analysis_years = int(deal["analysis_years"])
    if analysis_years < 1:
        raise ValueError("analysis_years must be >= 1")

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    refs = build_assumptions_tab(wb, pair, pair_verdict, deal_assumptions)
    refs = build_sources_uses_tab(wb, refs)
    refs = build_ppa_tab(wb, refs, pair_verdict)
    refs = build_proforma_tab(wb, refs, analysis_years)
    refs = build_accretion_tab(wb, refs, analysis_years)
    refs = build_concentration_tab(wb, refs, pair)
    refs = build_sensitivity_tab(wb, refs, deal)

    # Company identity metadata for downstream consumers (report generator,
    # comparison tool) that work from the .xlsx file alone
    for role, key in (("acquirer", "acq"), ("target", "tgt")):
        profile = pair.get(role) or {}
        refs[f"{key}_name"] = profile.get("company_name")
        refs[f"{key}_ticker"] = profile.get("ticker")
        refs[f"{key}_sector"] = (f"{profile.get('sic_code', '')} - "
                                 f"{profile.get('sic_description', '')}")

    _store_refs_in_workbook(wb, refs)

    wb.save(output_path)
    return refs


# =============================================================================
# REFS PERSISTENCE (hidden sheet so consumers can work from the file alone)
# =============================================================================

_REFS_SHEET = "_refs"


def _refs_to_json_safe(refs: dict) -> dict:
    """Deep-copy refs with the tuple-keyed sensitivity grid made JSON-safe."""
    import copy
    out = copy.deepcopy(refs)
    grid = out.get("sensitivity_grid")
    if grid and isinstance(grid.get("cells"), dict):
        grid["cells"] = [[p, s, ref] for (p, s), ref in grid["cells"].items()]
    return out


def _refs_from_json_safe(raw: dict) -> dict:
    """Inverse of _refs_to_json_safe."""
    grid = raw.get("sensitivity_grid")
    if grid and isinstance(grid.get("cells"), list):
        grid["cells"] = {(p, s): ref for p, s, ref in grid["cells"]}
    return raw


def _store_refs_in_workbook(wb: Workbook, refs: dict) -> None:
    import json
    ws = wb.create_sheet(_REFS_SHEET)
    ws.sheet_state = "hidden"
    ws["A1"] = json.dumps(_refs_to_json_safe(refs))


def load_refs_from_workbook(path: str) -> dict:
    """Load the refs map embedded in a generated workbook."""
    import json
    from openpyxl import load_workbook as _load
    wb = _load(path, data_only=False)
    if _REFS_SHEET not in wb.sheetnames:
        raise ValueError(
            f"No embedded refs found in {path} — was this workbook generated "
            f"by generate_ma_workbook?")
    raw = json.loads(wb[_REFS_SHEET]["A1"].value)
    return _refs_from_json_safe(raw)
