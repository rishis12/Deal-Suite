"""
M&A Modeling Tool - Report Generator

Generates a narrative M&A analysis report from a generated workbook,
adapted from AIO LBO's report_generator.py (/reference/backend/).

Carried over directly:
- Real recalculation before extraction (recalc.py: LibreOffice headless
  preferred, Excel COM fallback; read back with data_only=True — never
  trust unrecalculated formula strings)
- The BYOK multi-provider LLM abstraction (Anthropic, OpenAI, Gemini);
  API keys are never stored, logged, or written to disk, and provider
  errors never leak the key
- The provenance tagging system: fetched, defaulted, substituted,
  user_assumption, calculated, plus user_provided (Phase 5's market
  share table entries)
- The disclosure discipline: anything defaulted or substituted is stated
  plainly in the narrative, not buried

M&A-specific scope boundary (Section 6, Market Concentration): the
narrative states computed HHI figures and their standard DOJ/FTC
classification ONLY — it must never speculate about regulatory approval,
review timelines, remedies, or any real-world antitrust outcome.
"""

import re
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .recalc import recalculate_workbook


# =============================================================================
# PART 1: DATA EXTRACTION WITH PROVENANCE TAGGING
# =============================================================================

SOURCE_TYPES = ("fetched", "defaulted", "substituted", "user_assumption",
                "calculated", "user_provided")


@dataclass
class TaggedValue:
    """A value with explicit provenance tracking (carried over from AIO LBO)."""
    value: Any
    source: str
    note: Optional[str] = None

    def to_dict(self) -> dict:
        result = {"value": self.value, "source": self.source}
        if self.note:
            result["note"] = self.note
        return result


@dataclass
class MAExtractedData:
    """All extracted data from a recalculated M&A workbook, tagged."""
    values: Dict[str, TaggedValue] = field(default_factory=dict)

    # Year-by-year series (calculated values)
    proforma_eps_by_year: List[float] = field(default_factory=list)
    ad_by_year: List[float] = field(default_factory=list)

    # Market concentration state
    hhi_assessed: bool = False
    pre_merger_hhi: Optional[float] = None
    post_merger_hhi: Optional[float] = None
    delta_hhi: Optional[float] = None
    pre_merger_class: Optional[str] = None
    post_merger_class: Optional[str] = None
    presumptive_concern: Optional[str] = None

    # Accretion/dilution crossover (computed deterministically in Python,
    # handed to the LLM as a fact — never left for the LLM to derive)
    crossover: Optional[str] = None

    def get(self, key: str) -> Any:
        tv = self.values.get(key)
        return tv.value if tv else None

    def get_defaulted_fields(self) -> List[Tuple[str, str]]:
        return [(k, tv.note or "defaulted") for k, tv in self.values.items()
                if tv.source in ("defaulted", "substituted")]

    def get_user_provided_fields(self) -> List[Tuple[str, Any, str]]:
        return [(k, tv.value, tv.note or "Manually provided by user")
                for k, tv in self.values.items() if tv.source == "user_provided"]

    def to_dict(self) -> dict:
        return {
            "values": {k: tv.to_dict() for k, tv in self.values.items()},
            "proforma_eps_by_year": self.proforma_eps_by_year,
            "ad_by_year": self.ad_by_year,
            "hhi_assessed": self.hhi_assessed,
            "pre_merger_hhi": self.pre_merger_hhi,
            "post_merger_hhi": self.post_merger_hhi,
            "delta_hhi": self.delta_hhi,
            "pre_merger_class": self.pre_merger_class,
            "post_merger_class": self.post_merger_class,
            "presumptive_concern": self.presumptive_concern,
            "crossover": self.crossover,
        }


_REF_RE = re.compile(r"'(.+)'!\$([A-Z]+)\$(\d+)")


def _read_ref(wb, ref: str):
    m = _REF_RE.match(ref)
    if not m:
        return None
    sheet, col, row = m.group(1), m.group(2), int(m.group(3))
    if sheet not in wb.sheetnames:
        return None
    return wb[sheet][f"{col}{row}"].value


def _cell_provenance(wb_fmt, ref: str, default_source: str) -> Tuple[str, Optional[str]]:
    """
    Provenance from cell styling in the ORIGINAL (formula) workbook, using
    the AIO LBO color convention:
    - Green fill (C6EFCE) = user_provided
    - Yellow fill (FFFF99) = defaulted
    - Orange fill (FFB366) = substituted (e.g. implausible growth fallback)
    - otherwise the caller's default (fetched / user_assumption / calculated)
    """
    m = _REF_RE.match(ref)
    if not m:
        return default_source, None
    cell = wb_fmt[m.group(1)][f"{m.group(2)}{int(m.group(3))}"]
    comment = cell.comment.text if cell.comment else None

    fill = cell.fill
    if fill and fill.start_color and fill.start_color.rgb:
        rgb = str(fill.start_color.rgb).upper()
        if 'C6EFCE' in rgb:
            return "user_provided", comment
        if 'FFFF99' in rgb:
            return "defaulted", comment
        if 'FFB366' in rgb:
            return "substituted", comment
    return default_source, None


def extract_data_from_workbook(recalc_path: str, original_path: str,
                               refs: dict) -> MAExtractedData:
    """
    Extract all report-relevant data from a RECALCULATED workbook, using
    the refs map produced by generate_ma_workbook, with provenance from
    the original workbook's cell styling.
    """
    wb = load_workbook(recalc_path, data_only=True)
    wb_fmt = load_workbook(original_path, data_only=False)

    data = MAExtractedData()

    def tag(key: str, ref_key: str, default_source: str):
        ref = refs.get(ref_key)
        if not ref or not isinstance(ref, str):
            return
        value = _read_ref(wb, ref)
        source, note = _cell_provenance(wb_fmt, ref, default_source)
        data.values[key] = TaggedValue(value=value, source=source, note=note)

    # Company data (fetched)
    for role in ("acq", "tgt"):
        tag(f"{role}_price", f"{role}_price", "fetched")
        tag(f"{role}_diluted_shares", f"{role}_diluted_shares", "fetched")
        tag(f"{role}_diluted_eps", f"{role}_diluted_eps", "fetched")
        tag(f"{role}_revenue", f"{role}_revenue", "fetched")
        tag(f"{role}_op_income", f"{role}_op_income", "fetched")
        tag(f"{role}_da", f"{role}_da", "fetched")
        tag(f"{role}_net_income", f"{role}_net_income", "fetched")
        tag(f"{role}_total_debt", f"{role}_total_debt", "fetched")
        tag(f"{role}_cash", f"{role}_cash", "fetched")
        tag(f"{role}_total_assets", f"{role}_total_assets", "fetched")
        tag(f"{role}_total_liabilities", f"{role}_total_liabilities", "fetched")
        tag(f"{role}_growth", f"{role}_growth", "user_assumption")

    # Deal assumptions (user_assumption)
    for key in ("offer_premium", "pct_cash", "pct_debt", "new_debt_rate",
                "foregone_cash_rate", "revenue_synergies_pct",
                "cost_synergies_pct", "synergy_ramp_years",
                "transaction_fee_pct", "asset_step_up_pct",
                "intangible_amort_years", "tax_rate", "analysis_years"):
        tag(key, key, "user_assumption")
    tag("pct_stock", "pct_stock", "calculated")

    # Sources & Uses / PPA (calculated)
    for key in ("purchase_equity_value", "transaction_fees", "total_uses",
                "cash_used", "new_debt", "new_stock_issued",
                "new_shares_issued", "balance_check", "tgt_book_equity",
                "premium_over_book", "asset_step_up", "goodwill",
                "incremental_da", "standalone_eps"):
        tag(key, key, "calculated")

    # Company names/tickers/sectors come through refs-independent metadata:
    # stored by the caller on the refs dict (see generate_report / API).
    for key in ("acq_name", "tgt_name", "acq_ticker", "tgt_ticker",
                "acq_sector", "tgt_sector"):
        if key in refs:
            data.values[key] = TaggedValue(value=refs[key], source="fetched")

    # Pro forma trajectory (year 1 and final year)
    pf = refs.get("proforma_rows", {})
    analysis_years = int(data.get("analysis_years") or 5)
    if pf:
        def pf_cell(row_key, year):
            col = chr(66 + year)
            return wb["Pro Forma IS"][f"{col}{pf[row_key]}"].value

        for label, row_key in (("comb_revenue", "comb_rev_row"),
                               ("comb_ebit", "comb_ebit_row"),
                               ("comb_net_income", "comb_ni_row"),
                               ("rev_synergies", "rev_syn_row"),
                               ("cost_synergies", "cost_syn_row")):
            data.values[f"{label}_year1"] = TaggedValue(
                value=pf_cell(row_key, 1), source="calculated")
            data.values[f"{label}_final"] = TaggedValue(
                value=pf_cell(row_key, analysis_years), source="calculated")

    # EPS + A/D series
    for ref in refs.get("proforma_eps_cells", []):
        data.proforma_eps_by_year.append(_read_ref(wb, ref))
    for ref in refs.get("ad_cells", []):
        data.ad_by_year.append(_read_ref(wb, ref))

    # Crossover detection (deterministic)
    ad = [v for v in data.ad_by_year]
    if ad and all(v is not None for v in ad):
        signs = ["accretive" if v > 0 else "dilutive" if v < 0 else "neutral"
                 for v in ad]
        first = signs[0]
        crossover = None
        for i, s in enumerate(signs[1:], start=2):
            if s != first and s != "neutral":
                crossover = (f"The deal is {first} in Year 1 and crosses to "
                             f"{s} in Year {i}.")
                break
        if crossover is None:
            crossover = f"The deal stays {first} across all {len(ad)} years."
        data.crossover = crossover

    # Market concentration
    pre = _read_ref(wb, refs.get("pre_merger_hhi", ""))
    if isinstance(pre, (int, float)):
        data.hhi_assessed = True
        data.pre_merger_hhi = float(pre)
        post = _read_ref(wb, refs.get("post_merger_hhi", ""))
        delta = _read_ref(wb, refs.get("delta_hhi", ""))
        data.post_merger_hhi = float(post) if isinstance(post, (int, float)) else None
        data.delta_hhi = float(delta) if isinstance(delta, (int, float)) else None
        data.pre_merger_class = _read_ref(wb, refs.get("pre_merger_class", ""))
        data.post_merger_class = _read_ref(wb, refs.get("post_merger_class", ""))
        data.presumptive_concern = _read_ref(wb, refs.get("presumptive_concern", ""))
        # Market share entries are user_provided by definition (Phase 5)
        table = refs.get("market_share_table", {})
        sheet = table.get("sheet")
        if sheet:
            shares = []
            for name_addr, share_addr in zip(table.get("name_cells", []),
                                             table.get("share_cells", [])):
                name = wb[sheet][name_addr].value
                share = wb[sheet][share_addr].value
                if share is not None:
                    shares.append((name, share))
            if shares:
                data.values["market_shares"] = TaggedValue(
                    value=shares, source="user_provided",
                    note="Market share estimates entered by the user "
                         "(externally researched, human-reviewed)")
    else:
        data.hhi_assessed = False

    return data


# =============================================================================
# PART 2: LLM PROVIDER ABSTRACTION (BYOK — carried over from AIO LBO)
# =============================================================================

class LLMProviderError(Exception):
    """Error from an LLM provider that does NOT leak the API key."""
    pass


def _call_anthropic(prompt: str, api_key: str) -> str:
    """Call Anthropic Claude API. Raises LLMProviderError without leaking the key."""
    try:
        import anthropic
    except ImportError:
        raise LLMProviderError(
            "anthropic package not installed. Run: pip install anthropic"
        )

    try:
        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model="claude-sonnet-4-20250514",
            max_tokens=2000,
            messages=[{"role": "user", "content": prompt}]
        )
        return response.content[0].text

    except anthropic.AuthenticationError:
        raise LLMProviderError("Anthropic API authentication failed. Check your API key.")
    except anthropic.RateLimitError:
        raise LLMProviderError("Anthropic API rate limit exceeded. Try again later.")
    except anthropic.APIStatusError as e:
        raise LLMProviderError(f"Anthropic API error: {e.message}")
    except Exception as e:
        raise LLMProviderError(f"Anthropic API call failed: {type(e).__name__}")


def _call_openai(prompt: str, api_key: str) -> str:
    """Call OpenAI GPT API. Raises LLMProviderError without leaking the key."""
    try:
        import openai
    except ImportError:
        raise LLMProviderError(
            "openai package not installed. Run: pip install openai"
        )

    try:
        client = openai.OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model="gpt-4o",
            max_tokens=2000,
            messages=[{"role": "user", "content": prompt}]
        )
        return response.choices[0].message.content

    except openai.AuthenticationError:
        raise LLMProviderError("OpenAI API authentication failed. Check your API key.")
    except openai.RateLimitError:
        raise LLMProviderError("OpenAI API rate limit exceeded. Try again later.")
    except openai.APIStatusError as e:
        raise LLMProviderError(f"OpenAI API error: {e.message}")
    except Exception as e:
        raise LLMProviderError(f"OpenAI API call failed: {type(e).__name__}")


def _call_gemini(prompt: str, api_key: str) -> str:
    """Call Google Gemini API. Raises LLMProviderError without leaking the key."""
    try:
        from google import genai
        from google.genai import types
    except ImportError:
        raise LLMProviderError(
            "google-genai package not installed. Run: pip install google-genai"
        )

    try:
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt,
            config=types.GenerateContentConfig(max_output_tokens=4000)
        )
        return response.text

    except Exception as e:
        error_type = type(e).__name__
        error_msg = str(e).lower()
        if 'api_key' in error_msg or 'authentication' in error_msg or 'invalid' in error_msg:
            raise LLMProviderError("Gemini API authentication failed. Check your API key.")
        if 'quota' in error_msg or 'rate' in error_msg or 'limit' in error_msg:
            raise LLMProviderError("Gemini API rate limit or quota exceeded. Try again later.")
        if 'blocked' in error_msg or 'safety' in error_msg:
            raise LLMProviderError("Gemini API blocked the request due to safety filters.")
        if '404' in error_msg or 'not found' in error_msg:
            raise LLMProviderError("Gemini model not found. The model name may have changed.")
        raise LLMProviderError(f"Gemini API call failed: {error_type}")


# Provider registry - easy to add more providers
LLM_PROVIDERS: Dict[str, Callable[[str, str], str]] = {
    "anthropic": _call_anthropic,
    "openai": _call_openai,
    "gemini": _call_gemini,
}


def generate_narrative(data: MAExtractedData, provider: str, api_key: str,
                       prompt_override: Optional[str] = None) -> str:
    """
    Generate the narrative report text via the chosen BYOK provider.
    The api_key is not stored, logged, or written to disk.
    """
    if provider not in LLM_PROVIDERS:
        raise ValueError(
            f"Unsupported LLM provider: {provider}. "
            f"Supported: {list(LLM_PROVIDERS.keys())}")

    prompt = prompt_override if prompt_override else build_report_prompt(data)
    return LLM_PROVIDERS[provider](prompt, api_key)


# =============================================================================
# PART 3: THE PROMPT
# =============================================================================

def format_currency(value) -> str:
    if value is None:
        return "N/A"
    if abs(value) >= 1e9:
        return f"${value/1e9:,.2f}B"
    if abs(value) >= 1e6:
        return f"${value/1e6:,.2f}M"
    return f"${value:,.0f}"


def format_pct(value) -> str:
    if value is None:
        return "N/A"
    return f"{value*100:.1f}%"


def build_report_prompt(data: MAExtractedData) -> str:
    """Build the LLM prompt for the M&A narrative report."""
    g = data.get

    ad_series = ", ".join(
        f"Year {i+1}: {format_pct(v)}" for i, v in enumerate(data.ad_by_year))

    prompt = f"""You are writing a concise M&A accretion/dilution analysis report. Use ONLY the data provided below. Do not invent, estimate, or speculate about anything not explicitly stated.

=== DEAL ===
Acquirer: {g('acq_name') or 'Unknown'} ({g('acq_ticker') or '?'}) — {g('acq_sector') or 'Unknown sector'}
Target: {g('tgt_name') or 'Unknown'} ({g('tgt_ticker') or '?'}) — {g('tgt_sector') or 'Unknown sector'}
Offer Premium: {format_pct(g('offer_premium'))}
Financing Mix: {format_pct(g('pct_cash'))} cash / {format_pct(g('pct_debt'))} debt / {format_pct(g('pct_stock'))} stock
New Debt Interest Rate: {format_pct(g('new_debt_rate'))}
Foregone Interest Rate on Cash: {format_pct(g('foregone_cash_rate'))}
Revenue Synergies: {format_pct(g('revenue_synergies_pct'))} of combined revenue
Cost Synergies: {format_pct(g('cost_synergies_pct'))} of combined opex
Synergy Ramp: {g('synergy_ramp_years')} years to full run-rate
Tax Rate: {format_pct(g('tax_rate'))}
Analysis Horizon: {g('analysis_years')} years

=== SOURCES & USES ===
Purchase Equity Value: {format_currency(g('purchase_equity_value'))}
Transaction Fees: {format_currency(g('transaction_fees'))}
Total Uses: {format_currency(g('total_uses'))}
Cash Used: {format_currency(g('cash_used'))}
New Debt Raised: {format_currency(g('new_debt'))}
New Stock Issued: {format_currency(g('new_stock_issued'))} ({f"{g('new_shares_issued'):,.0f}" if g('new_shares_issued') is not None else 'N/A'} new shares)

=== PURCHASE PRICE ALLOCATION ===
Target Book Equity: {format_currency(g('tgt_book_equity'))}
Premium Over Book: {format_currency(g('premium_over_book'))}
Asset Step-Up: {format_currency(g('asset_step_up'))}
Goodwill: {format_currency(g('goodwill'))}
Incremental Annual D&A: {format_currency(g('incremental_da'))}

=== PRO FORMA TRAJECTORY ===
Combined Revenue: Year 1 {format_currency(g('comb_revenue_year1'))} -> Final Year {format_currency(g('comb_revenue_final'))}
Combined EBIT: Year 1 {format_currency(g('comb_ebit_year1'))} -> Final Year {format_currency(g('comb_ebit_final'))}
Revenue Synergies Realized: Year 1 {format_currency(g('rev_synergies_year1'))} -> Final Year {format_currency(g('rev_synergies_final'))}
Cost Synergies Realized: Year 1 {format_currency(g('cost_synergies_year1'))} -> Final Year {format_currency(g('cost_synergies_final'))}
Combined Net Income: Year 1 {format_currency(g('comb_net_income_year1'))} -> Final Year {format_currency(g('comb_net_income_final'))}
Acquirer Revenue Growth Assumption: {format_pct(g('acq_growth'))}
Target Revenue Growth Assumption: {format_pct(g('tgt_growth'))}

=== ACCRETION / DILUTION (THE HEADLINE) ===
Acquirer Standalone EPS: {f"${g('standalone_eps'):,.2f}" if g('standalone_eps') is not None else 'N/A'}
Accretion/(Dilution) by year: {ad_series}
Year 1: {format_pct(data.ad_by_year[0]) if data.ad_by_year else 'N/A'}
Final Year: {format_pct(data.ad_by_year[-1]) if data.ad_by_year else 'N/A'}
CROSSOVER FACT (computed, state this explicitly in the report): {data.crossover or 'N/A'}

"""

    # Market concentration block
    if data.hhi_assessed:
        shares_tv = data.values.get("market_shares")
        shares_str = ""
        if shares_tv:
            shares_str = "; ".join(
                f"{name or 'Unnamed'}: {format_pct(share)}"
                for name, share in shares_tv.value)
        prompt += f"""=== MARKET CONCENTRATION (user-provided market share estimates) ===
Market shares entered by the user (user_provided; externally researched, human-reviewed): {shares_str}
Pre-Merger HHI: {data.pre_merger_hhi:,.0f} ({data.pre_merger_class})
Post-Merger HHI: {data.post_merger_hhi:,.0f} ({data.post_merger_class})
Delta HHI: {data.delta_hhi:,.0f}
Presumptive-Concern Trigger (Delta > 200 AND Post > 2,500): {data.presumptive_concern}

"""
    else:
        prompt += """=== MARKET CONCENTRATION ===
NOT ASSESSED — the market share table was left empty, so no HHI was calculated.

"""

    # Disclosures
    user_provided_fields = data.get_user_provided_fields()
    if user_provided_fields:
        prompt += "=== USER-PROVIDED DATA (MUST be mentioned in report) ===\n"
        prompt += ("The following values were manually provided by the user "
                   "rather than fetched from data sources:\n")
        for field_name, value, note in user_provided_fields:
            if field_name == "market_shares":
                prompt += ("- Market share estimates in the concentration "
                           "analysis were entered by the user\n")
            else:
                prompt += f"- {field_name}: {value} ({note})\n"
        prompt += "\n"

    defaulted_fields = data.get_defaulted_fields()
    if defaulted_fields:
        prompt += "=== DATA QUALITY DISCLOSURES (MUST be mentioned in report) ===\n"
        prompt += "The following values were defaulted or substituted due to missing data:\n"
        for field_name, note in defaulted_fields:
            prompt += f"- {field_name}: {note}\n"
        prompt += "\n"

    prompt += """=== YOUR TASK ===

Write a concise M&A analysis report with the following sections:

1. **Deal Summary** (one paragraph): Who is acquiring whom, the offer premium, and the cash/debt/stock financing mix.

2. **Sources & Uses** (brief, factual): Total transaction cost and how it is funded.

3. **Purchase Price Allocation** (brief): Goodwill and asset step-up created, and the incremental annual D&A.

4. **Pro Forma Trajectory**: How combined revenue and EBIT evolve over the analysis horizon and how synergies ramp in.

5. **Accretion/Dilution — THE HEADLINE SECTION**: State the Year 1 and final-year accretion/(dilution) figures clearly. You are given a computed CROSSOVER FACT above — state it explicitly. If the deal crosses from dilutive to accretive (or vice versa) partway through the period, that crossover is the single most important fact in this report and must not get lost.

6. **Market Concentration — CRITICAL SCOPE BOUNDARY**: If HHI figures are provided above, state the computed HHI figures and their standard classification (e.g. "Post-merger HHI of X classifies as Highly Concentrated per DOJ/FTC guidelines, with a Delta HHI of Y"). State the computed HHI figures and their standard classification. Do NOT predict, speculate on, or comment on the likelihood of regulatory approval, required divestitures, review timelines, or any other real-world antitrust outcome — you have no information about any of that. If the section above says NOT ASSESSED, state plainly that market concentration was not assessed because no market share estimates were provided — do not invent commentary about it.

7. **Verdict**: A short, plain, FINANCIAL-ONLY statement: does the deal look accretive or dilutive and by how much. This is explicitly NOT a recommendation on whether to proceed with the transaction — do not phrase it as one.

IMPORTANT INSTRUCTIONS:
- Stay factual and grounded in the numbers provided. Do not invent, estimate, or speculate about anything not in the data.
- Any field listed under "USER-PROVIDED DATA" above MUST be explicitly disclosed as user-entered.
- Any field listed under "DATA QUALITY DISCLOSURES" above MUST be explicitly disclosed in plain language (e.g. "Total debt was not reported and was assumed to be $0."). Do not bury or soften these.
- Keep it concise: a few short paragraphs total, not an exhaustive essay.
- No investment advice framing. Describe what the numbers show.
- Format currency values consistently ($X.XXB for billions, $X.XXM for millions).
"""

    return prompt


# =============================================================================
# MAIN REPORT GENERATION FUNCTION
# =============================================================================

def generate_report(xlsx_path: str, refs: dict, provider: str, api_key: str,
                    pair: Optional[dict] = None) -> Tuple[str, MAExtractedData]:
    """
    Generate a complete M&A narrative report from a generated workbook.

    Args:
        xlsx_path: path to the generated (unrecalculated) .xlsx
        refs: the refs dict returned by generate_ma_workbook
        provider: "anthropic" | "openai" | "gemini"
        api_key: BYOK key (never stored, logged, or written to disk)
        pair: optional fetch_ma_pair result to supply company names/tickers

    Returns (report_text, extracted_data).
    """
    if pair:
        refs = dict(refs)
        for role, key in (("acquirer", "acq"), ("target", "tgt")):
            profile = pair.get(role) or {}
            refs[f"{key}_name"] = profile.get("company_name")
            refs[f"{key}_ticker"] = profile.get("ticker")
            refs[f"{key}_sector"] = (f"{profile.get('sic_code', '')} - "
                                     f"{profile.get('sic_description', '')}")

    print("Recalculating workbook before extraction...")
    recalc_path = recalculate_workbook(xlsx_path)

    print("Extracting data with provenance tags...")
    data = extract_data_from_workbook(recalc_path, xlsx_path, refs)

    print(f"Generating narrative with {provider}...")
    report = generate_narrative(data, provider, api_key)

    return report, data
