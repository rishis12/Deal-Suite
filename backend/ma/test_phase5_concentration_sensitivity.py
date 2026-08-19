"""
Phase 5 test: Market Concentration (HHI) + Sensitivity tabs.

1. Hand-verified HHI with a constructed 40/35/25 example (known answer:
   pre=3450, post=6250, delta=2800) before any real-data complexity
2. Threshold classification at boundary values (just above/below 1500 and
   2500, delta just above/below 200)
3. CRITICAL: sensitivity grid's base-case cell must EXACTLY match the
   Accretion/Dilution tab's Year-1 value after real recalculation
4. Grid monotonicity: strictly worse as premium rises (each column), and
   each row monotonic in % stock

All values read from REAL recalculated workbooks (data_only=True).
"""

import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

from . import data_layer as dl
from .validator import validate_pair
from .excel_generator import generate_ma_workbook
from .recalc import recalculate_workbook

OUT_DIR = Path(__file__).parent / "test_output"
failures = []


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    print(f"  [{status}] {name}" + (f" — {detail}" if detail else ""))
    if not condition:
        failures.append(f"{name}: {detail}")


def read_ref(wb, ref: str):
    m = re.match(r"'(.+)'!\$([A-Z]+)\$(\d+)", ref)
    return wb[m.group(1)][f"{m.group(2)}{int(m.group(3))}"].value


def inject_shares_and_recalc(base_path, refs, shares: list, tag: str):
    """Copy workbook, write share values into the user table, recalc."""
    variant = str(OUT_DIR / f"p5_{tag}.xlsx")
    shutil.copy(base_path, variant)
    wb = load_workbook(variant)
    ws = wb["Market Concentration"]
    for cell_addr, value in zip(refs["market_share_table"]["share_cells"], shares):
        ws[cell_addr] = value
    wb.save(variant)
    return load_workbook(recalculate_workbook(variant), data_only=True)


def main():
    OUT_DIR.mkdir(exist_ok=True)

    print("Fetching MSFT/POOL and generating workbook...")
    pair = dl.fetch_ma_pair("MSFT", "POOL", verbose=False, skip_price=True)
    pair["acquirer"]["current_price"] = 500.0
    pair["target"]["current_price"] = 320.0
    for role in ("acquirer", "target"):
        pair[role]["current_price_source"] = "synthetic-test"
    verdict = validate_pair(pair)
    assert verdict["status"] != "fail"

    base_path = str(OUT_DIR / "p5_base_MSFT_POOL.xlsx")
    refs = generate_ma_workbook(pair, verdict, base_path)

    # ==================================================================
    # 0. Empty-table state: no error, no misleading 0
    # ==================================================================
    print(f"\n{'='*70}\n0. EMPTY-TABLE STATE\n{'='*70}")
    wb0 = load_workbook(recalculate_workbook(base_path), data_only=True)
    pre0 = read_ref(wb0, refs["pre_merger_hhi"])
    check("empty table shows explicit not-yet-calculated message",
          isinstance(pre0, str) and "not yet calculated" in pre0, f"pre={pre0!r}")
    check("empty table classification shows placeholder",
          read_ref(wb0, refs["pre_merger_class"]) == "—",
          repr(read_ref(wb0, refs["pre_merger_class"])))

    # ==================================================================
    # 1. Hand-verified constructed HHI example: 40 / 35 / 25
    # ==================================================================
    print(f"\n{'='*70}\n1. CONSTRUCTED HHI: 40%/35%/25%\n{'='*70}")
    wb1 = inject_shares_and_recalc(base_path, refs, [0.40, 0.35, 0.25], "hhi_403525")
    pre = read_ref(wb1, refs["pre_merger_hhi"])
    post = read_ref(wb1, refs["post_merger_hhi"])
    delta = read_ref(wb1, refs["delta_hhi"])
    print(f"  pre={pre}, post={post}, delta={delta}")
    check("Pre-Merger HHI == 3450 (1600+1225+625)", abs(pre - 3450) < 1e-6, f"{pre!r}")
    check("Post-Merger HHI == 6250 ((40+35)^2 + 25^2)", abs(post - 6250) < 1e-6, f"{post!r}")
    check("Delta HHI == 2800", abs(delta - 2800) < 1e-6, f"{delta!r}")
    check("post classification == Highly Concentrated",
          read_ref(wb1, refs["post_merger_class"]) == "Highly Concentrated")
    check("presumptive concern triggered (delta 2800 > 200, post > 2500)",
          "PRESUMED" in read_ref(wb1, refs["presumptive_concern"]))

    # ==================================================================
    # 2. Threshold boundaries
    # ==================================================================
    print(f"\n{'='*70}\n2. THRESHOLD BOUNDARIES\n{'='*70}")
    # Single-company share s: pre-HHI = s^2 * 10000 (target row left blank)
    boundary_cases = [
        # (acq share, expected pre-HHI approx, expected classification)
        (0.3872, 1499.2, "Unconcentrated"),          # just below 1500
        (0.3875, 1501.6, "Moderately Concentrated"), # just above 1500
        (0.5000, 2500.0, "Moderately Concentrated"), # exactly 2500 (<=)
        (0.5002, 2502.0, "Highly Concentrated"),     # just above 2500
    ]
    for share, exp_hhi, exp_class in boundary_cases:
        wbb = inject_shares_and_recalc(base_path, refs, [share],
                                       f"hhi_bound_{int(share*10000)}")
        pre = read_ref(wbb, refs["pre_merger_hhi"])
        cls = read_ref(wbb, refs["pre_merger_class"])
        check(f"HHI {pre:.1f} -> {exp_class}",
              cls == exp_class and abs(pre - exp_hhi) < 1.0,
              f"pre={pre!r}, class={cls!r}")

    # Delta boundary: delta = 2*a*t*10000. post > 2500 in both cases.
    # a=50%, t=2.0% -> delta = 200 exactly (NOT > 200), post = 2704 -> below
    # a=50%, t=2.1% -> delta = 210 -> triggered (post = 2714 > 2500)
    wbd = inject_shares_and_recalc(base_path, refs, [0.50, 0.02], "delta_200")
    delta = read_ref(wbd, refs["delta_hhi"])
    concern = read_ref(wbd, refs["presumptive_concern"])
    check(f"delta exactly 200 does NOT trigger (delta={delta:.2f})",
          abs(delta - 200) < 0.5 and "Below" in concern, f"{concern!r}")

    wbd2 = inject_shares_and_recalc(base_path, refs, [0.50, 0.021], "delta_220")
    delta2 = read_ref(wbd2, refs["delta_hhi"])
    post2 = read_ref(wbd2, refs["post_merger_hhi"])
    concern2 = read_ref(wbd2, refs["presumptive_concern"])
    check(f"delta {delta2:.1f} > 200 with post {post2:.1f} > 2500 triggers",
          "PRESUMED" in concern2, f"{concern2!r}")

    # And: delta > 200 but post BELOW 2500 must NOT trigger
    wbd3 = inject_shares_and_recalc(base_path, refs, [0.10, 0.15], "delta_lowpost")
    delta3 = read_ref(wbd3, refs["delta_hhi"])
    post3 = read_ref(wbd3, refs["post_merger_hhi"])
    concern3 = read_ref(wbd3, refs["presumptive_concern"])
    check(f"delta {delta3:.0f} > 200 but post {post3:.0f} < 2500 -> no trigger",
          delta3 > 200 and post3 < 2500 and "Below" in concern3, f"{concern3!r}")

    # ==================================================================
    # 3. CRITICAL: base-case grid cell == A/D tab Year 1, exactly
    # ==================================================================
    print(f"\n{'='*70}\n3. SENSITIVITY BASE CELL vs A/D TAB\n{'='*70}")
    grid = refs["sensitivity_grid"]
    base_cell_val = read_ref(wb0, grid["base_cell"])
    ad_y1 = read_ref(wb0, refs["ad_year1"])
    print(f"  grid base cell ({grid['base_premium']:.1%} premium, "
          f"{grid['base_stock']:.1%} stock) = {base_cell_val!r}")
    print(f"  A/D tab Year 1 = {ad_y1!r}")
    check("grid base cell EXACTLY matches A/D tab Year 1",
          base_cell_val is not None and ad_y1 is not None
          and abs(base_cell_val - ad_y1) < 1e-9,
          f"grid={base_cell_val!r} vs ad={ad_y1!r}")

    # ==================================================================
    # 4. Monotonicity
    # ==================================================================
    print(f"\n{'='*70}\n4. GRID MONOTONICITY\n{'='*70}")
    values = {}
    for (p, s), ref in grid["cells"].items():
        values[(p, s)] = read_ref(wb0, ref)

    # Higher premium must always be worse (strictly decreasing per column):
    # a higher price raises financing costs and share issuance with no
    # offsetting benefit.
    ok = True
    for s in grid["stock_mixes"]:
        col_vals = [values[(p, s)] for p in grid["premiums"]]
        if not all(col_vals[i] > col_vals[i+1] for i in range(len(col_vals)-1)):
            ok = False
            print(f"    NOT decreasing in premium at stock={s}: {col_vals}")
    check("A/D strictly decreasing as premium rises (every column)", ok)

    # Each row must be monotonic in % stock (EPS is a ratio of two linear
    # functions of s, so its direction cannot flip mid-row).
    ok = True
    directions = []
    for p in grid["premiums"]:
        row_vals = [values[(p, s)] for s in grid["stock_mixes"]]
        inc = all(row_vals[i] <= row_vals[i+1] + 1e-12 for i in range(len(row_vals)-1))
        dec = all(row_vals[i] >= row_vals[i+1] - 1e-12 for i in range(len(row_vals)-1))
        directions.append("inc" if inc else "dec" if dec else "MIXED")
        if not (inc or dec):
            ok = False
            print(f"    non-monotonic row at premium={p}: {row_vals}")
    check("each row monotonic in % stock", ok, f"directions={directions}")

    # Direction sanity for THIS pair: MSFT at $500 with EPS ~11.8 has an
    # earnings yield of ~2.4%; the blended after-tax financing cost of
    # cash+debt (~8%*3/7 + 4.5%*4/7, taxed) is higher than the earnings
    # yield given POOL's earnings acquired, so MORE STOCK should generally
    # HELP accretion here (stock is the cheaper currency in EPS terms).
    row_base = [values[(grid["base_premium"], s)] for s in grid["stock_mixes"]]
    print(f"  base-premium row across stock mixes: "
          f"{[f'{v*100:+.1f}%' for v in row_base]}")
    check("more stock helps for this construction (base-premium row increasing)",
          row_base[0] < row_base[-1], f"{row_base[0]!r} -> {row_base[-1]!r}")

    print(f"\n{'='*70}")
    if failures:
        print(f"PHASE 5 TEST RESULT: {len(failures)} FAILURE(S)")
        for f in failures:
            print(f"  - {f}")
    else:
        print("PHASE 5 TEST RESULT: ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
