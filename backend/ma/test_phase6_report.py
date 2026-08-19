"""
Phase 6 test: report generator (extraction, provenance, prompt, BYOK plumbing).

No BYOK API key is available in this test environment, so the live LLM
narrative is exercised separately (the phase tester feeds the built prompt
to a real LLM and inspects the output). This script verifies everything
deterministic:

1. Extraction numbers trace exactly to the recalculated workbook
2. HHI block present when the market share table is filled; explicit
   NOT ASSESSED when empty
3. The regulatory scope-boundary instruction is present verbatim
4. Defaulted/substituted fields produce plain-language disclosure blocks
5. Crossover fact computed correctly against the actual A/D sign pattern
6. BYOK plumbing: mock provider round-trip, unknown provider rejected

Built prompts are saved to test_output/ for the phase tester's LLM run.
"""

import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

from . import data_layer as dl
from .validator import validate_pair
from .excel_generator import generate_ma_workbook
from .recalc import recalculate_workbook
from . import report_generator as rg

OUT_DIR = Path(__file__).parent / "test_output"
failures = []


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    print(f"  [{status}] {name}" + (f" — {detail}" if detail else ""))
    if not condition:
        failures.append(f"{name}: {detail}")


def read_ref(wb, ref):
    m = re.match(r"'(.+)'!\$([A-Z]+)\$(\d+)", ref)
    return wb[m.group(1)][f"{m.group(2)}{int(m.group(3))}"].value


def prep_pair(acq, tgt, acq_price, tgt_price):
    pair = dl.fetch_ma_pair(acq, tgt, verbose=False, skip_price=True)
    pair["acquirer"]["current_price"] = acq_price
    pair["target"]["current_price"] = tgt_price
    for role in ("acquirer", "target"):
        pair[role]["current_price_source"] = "synthetic-test"
    return pair


def extract(xlsx_path, refs, pair):
    """Mimic generate_report's extraction step without an LLM call."""
    refs = dict(refs)
    for role, key in (("acquirer", "acq"), ("target", "tgt")):
        profile = pair.get(role) or {}
        refs[f"{key}_name"] = profile.get("company_name")
        refs[f"{key}_ticker"] = profile.get("ticker")
        refs[f"{key}_sector"] = (f"{profile.get('sic_code', '')} - "
                                 f"{profile.get('sic_description', '')}")
    recalc_path = recalculate_workbook(xlsx_path)
    data = rg.extract_data_from_workbook(recalc_path, xlsx_path, refs)
    return data, recalc_path


def main():
    OUT_DIR.mkdir(exist_ok=True)

    # ==================================================================
    # CASE A: MSFT/POOL with market share table FILLED, crossover deal
    # (45% premium all-stock is dilutive Y1; 2.5% cost synergies ramping
    # over 3 years should pull it accretive later)
    # ==================================================================
    print(f"\n{'='*70}\nCASE A: filled market-share table + crossover\n{'='*70}")
    pair_a = prep_pair("MSFT", "POOL", 600.0, 100.0)
    verdict_a = validate_pair(pair_a)
    path_a = str(OUT_DIR / "p6_case_a.xlsx")
    refs_a = generate_ma_workbook(
        pair_a, verdict_a, path_a,
        deal_assumptions={"pct_cash": 0.0, "pct_debt": 0.0,
                          "offer_premium": 0.45, "cost_synergies_pct": 0.025})
    # Fill the market share table: MSFT 40%, POOL 25%, competitor 20%
    wb = load_workbook(path_a)
    ws = wb["Market Concentration"]
    for addr, val in zip(refs_a["market_share_table"]["share_cells"],
                         [0.40, 0.25, 0.20]):
        ws[addr] = val
    wb.save(path_a)

    data_a, recalc_a = extract(path_a, refs_a, pair_a)
    wb_a = load_workbook(recalc_a, data_only=True)

    # 1. Extraction traces to the recalculated workbook exactly
    trace_keys = ["purchase_equity_value", "total_uses", "goodwill",
                  "incremental_da", "new_shares_issued", "standalone_eps",
                  "tgt_book_equity", "cash_used", "new_debt"]
    for key in trace_keys:
        sheet_val = read_ref(wb_a, refs_a[key])
        check(f"extracted {key} == workbook value",
              data_a.get(key) == sheet_val,
              f"extracted={data_a.get(key)!r} vs sheet={sheet_val!r}")

    # A/D series traces
    ad_sheet = [read_ref(wb_a, r) for r in refs_a["ad_cells"]]
    check("A/D series matches workbook", data_a.ad_by_year == ad_sheet,
          f"{data_a.ad_by_year} vs {ad_sheet}")

    # 2. HHI assessed + correct numbers (hand: pre=40²+25²+20²=2625;
    #    post=(65)²+20²=4625; delta=2000)
    check("HHI assessed", data_a.hhi_assessed)
    check("pre HHI == 2625", data_a.pre_merger_hhi == 2625,
          repr(data_a.pre_merger_hhi))
    check("post HHI == 4625", data_a.post_merger_hhi == 4625,
          repr(data_a.post_merger_hhi))
    check("delta HHI == 2000", data_a.delta_hhi == 2000,
          repr(data_a.delta_hhi))
    check("presumptive concern triggered",
          "PRESUMED" in (data_a.presumptive_concern or ""))
    check("market shares tagged user_provided",
          data_a.values.get("market_shares") is not None
          and data_a.values["market_shares"].source == "user_provided")

    # 5. Crossover fact vs actual sign pattern
    signs = ["accretive" if v > 0 else "dilutive" for v in data_a.ad_by_year]
    print(f"  A/D pattern: {[f'{v*100:+.1f}%' for v in data_a.ad_by_year]}")
    print(f"  Crossover fact: {data_a.crossover}")
    if signs[0] == "dilutive" and "accretive" in signs:
        expected_year = signs.index("accretive") + 1
        check("crossover fact correct",
              data_a.crossover == f"The deal is dilutive in Year 1 and "
                                  f"crosses to accretive in Year {expected_year}.",
              data_a.crossover)
    else:
        check("constant-sign fact correct",
              data_a.crossover == f"The deal stays {signs[0]} across all {len(signs)} years.",
              f"{data_a.crossover} (pattern {signs})")

    # 3. Prompt contents
    prompt_a = rg.build_report_prompt(data_a)
    (OUT_DIR / "p6_prompt_case_a.txt").write_text(prompt_a, encoding="utf-8")
    check("scope-boundary instruction present verbatim",
          "Do NOT predict, speculate on, or comment on the likelihood of "
          "regulatory approval, required divestitures, review timelines, or "
          "any other real-world antitrust outcome" in prompt_a)
    check("HHI figures appear in prompt", "Pre-Merger HHI: 2,625" in prompt_a,
          "looked for 'Pre-Merger HHI: 2,625'")
    check("crossover fact appears in prompt",
          (data_a.crossover or "??") in prompt_a)
    check("prompt cites Year-1 A/D matching workbook",
          rg.format_pct(ad_sheet[0]) in prompt_a, rg.format_pct(ad_sheet[0]))

    # ==================================================================
    # CASE B: empty market share table
    # ==================================================================
    print(f"\n{'='*70}\nCASE B: empty market-share table\n{'='*70}")
    pair_b = prep_pair("AAPL", "CROX", 250.0, 90.0)
    verdict_b = validate_pair(pair_b)
    path_b = str(OUT_DIR / "p6_case_b.xlsx")
    refs_b = generate_ma_workbook(pair_b, verdict_b, path_b)
    data_b, _ = extract(path_b, refs_b, pair_b)

    check("HHI not assessed for empty table", not data_b.hhi_assessed)
    prompt_b = rg.build_report_prompt(data_b)
    (OUT_DIR / "p6_prompt_case_b.txt").write_text(prompt_b, encoding="utf-8")
    check("prompt says NOT ASSESSED plainly",
          "NOT ASSESSED — the market share table was left empty" in prompt_b)
    check("prompt instructs to state non-assessment, not invent",
          "state plainly that market concentration was not assessed" in prompt_b)

    # ==================================================================
    # CASE C: defaulted field disclosure (CCL/AAL — CCL genuinely lacks a
    # 'cash' value in some years; force-remove target total_debt to
    # deterministically exercise the yellow-fill defaulted path)
    # ==================================================================
    print(f"\n{'='*70}\nCASE C: defaulted-field disclosure\n{'='*70}")
    pair_c = prep_pair("CCL", "AAL", 25.0, 12.0)
    fy = max(pair_c["target"]["fiscal_years"])
    pair_c["target"]["fiscal_years"][fy]["total_debt"] = None  # forced missing
    verdict_c = validate_pair(pair_c)
    path_c = str(OUT_DIR / "p6_case_c.xlsx")
    refs_c = generate_ma_workbook(pair_c, verdict_c, path_c)
    data_c, _ = extract(path_c, refs_c, pair_c)

    defaulted = data_c.get_defaulted_fields()
    print(f"  defaulted/substituted fields: {[k for k, _ in defaulted]}")
    check("target total_debt tagged defaulted",
          any(k == "tgt_total_debt" for k, _ in defaulted), str(defaulted))
    check("CCL growth tagged substituted (COVID clamp)",
          data_c.values["acq_growth"].source == "substituted",
          data_c.values["acq_growth"].source)

    prompt_c = rg.build_report_prompt(data_c)
    (OUT_DIR / "p6_prompt_case_c.txt").write_text(prompt_c, encoding="utf-8")
    check("disclosure block present in prompt",
          "DATA QUALITY DISCLOSURES (MUST be mentioned in report)" in prompt_c)
    check("total_debt disclosure listed",
          "tgt_total_debt" in prompt_c)

    # ==================================================================
    # CASE D: genuine dilutive -> accretive crossover (JNJ buys WING at a
    # deliberately huge all-debt price; 2% cost synergies ramp over 3 yrs)
    # ==================================================================
    print(f"\n{'='*70}\nCASE D: genuine crossover deal\n{'='*70}")
    pair_d = prep_pair("JNJ", "WING", 160.0, 1800.0)
    verdict_d = validate_pair(pair_d)
    path_d = str(OUT_DIR / "p6_case_d.xlsx")
    refs_d = generate_ma_workbook(
        pair_d, verdict_d, path_d,
        deal_assumptions={"pct_cash": 0.0, "pct_debt": 1.0,
                          "new_debt_rate": 0.10, "cost_synergies_pct": 0.02})
    data_d, _ = extract(path_d, refs_d, pair_d)
    signs_d = ["accretive" if v > 0 else "dilutive" for v in data_d.ad_by_year]
    print(f"  A/D pattern: {[f'{v*100:+.1f}%' for v in data_d.ad_by_year]}")
    print(f"  Crossover fact: {data_d.crossover}")
    check("case D actually crosses sign",
          signs_d[0] == "dilutive" and "accretive" in signs_d, str(signs_d))
    if signs_d[0] == "dilutive" and "accretive" in signs_d:
        expected_year = signs_d.index("accretive") + 1
        check("crossover fact states the correct crossover year",
              data_d.crossover == f"The deal is dilutive in Year 1 and "
                                  f"crosses to accretive in Year {expected_year}.",
              data_d.crossover)
    prompt_d = rg.build_report_prompt(data_d)
    (OUT_DIR / "p6_prompt_case_d.txt").write_text(prompt_d, encoding="utf-8")
    check("crossover fact appears in case-D prompt",
          (data_d.crossover or "??") in prompt_d)

    # ==================================================================
    # 6. BYOK plumbing
    # ==================================================================
    print(f"\n{'='*70}\nBYOK PLUMBING\n{'='*70}")
    rg.LLM_PROVIDERS["mock"] = lambda prompt, key: f"MOCK-REPORT ({len(prompt)} chars)"
    try:
        out = rg.generate_narrative(data_b, "mock", "fake-key")
        check("mock provider round-trip", out.startswith("MOCK-REPORT"))
    finally:
        del rg.LLM_PROVIDERS["mock"]

    try:
        rg.generate_narrative(data_b, "nonexistent", "k")
        check("unknown provider rejected", False, "no exception raised")
    except ValueError:
        check("unknown provider rejected", True)

    # Providers are lazily imported but MUST be in requirements.txt
    reqs = (Path(__file__).parent / "requirements.txt").read_text()
    for pkg in ("anthropic", "openai", "google-genai"):
        check(f"{pkg} listed in requirements.txt", pkg in reqs)

    print(f"\n{'='*70}")
    if failures:
        print(f"PHASE 6 TEST RESULT: {len(failures)} FAILURE(S)")
        for f in failures:
            print(f"  - {f}")
    else:
        print("PHASE 6 TEST RESULT: ALL CHECKS PASSED")
        print(f"Prompts for live-LLM verification saved in {OUT_DIR}")


if __name__ == "__main__":
    main()
