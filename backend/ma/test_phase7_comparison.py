"""
Phase 7 test: comparison tool.

Mode A: same deal (MSFT acquiring POOL), all-cash vs all-stock, everything
else constant. Confirms the input diff contains ONLY the financing mix,
the output shift matches an independent first-principles recomputation
(direction reasoned before checking), and the commentary prompt applies
the single-variable causal exception plus the scope prohibition.

Mode B: two genuinely different deals with different market share tables.
Confirms side-by-side profiles, Deal A/Deal B framing, and the antitrust
scope prohibition in the prompt.
"""

from pathlib import Path

from openpyxl import load_workbook

from . import data_layer as dl
from .validator import validate_pair
from .excel_generator import generate_ma_workbook, load_refs_from_workbook
from . import comparison_tool as ct

OUT_DIR = Path(__file__).parent / "test_output"
failures = []


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    print(f"  [{status}] {name}" + (f" — {detail}" if detail else ""))
    if not condition:
        failures.append(f"{name}: {detail}")


def prep_pair(acq, tgt, acq_price, tgt_price):
    pair = dl.fetch_ma_pair(acq, tgt, verbose=False, skip_price=True)
    pair["acquirer"]["current_price"] = acq_price
    pair["target"]["current_price"] = tgt_price
    for role in ("acquirer", "target"):
        pair[role]["current_price_source"] = "synthetic-test"
    return pair


def fill_shares(path, refs, shares):
    wb = load_workbook(path)
    ws = wb["Market Concentration"]
    for addr, val in zip(refs["market_share_table"]["share_cells"], shares):
        ws[addr] = val
    wb.save(path)


def main():
    OUT_DIR.mkdir(exist_ok=True)

    # ==================================================================
    # Build the three workbooks
    # ==================================================================
    print("Building workbooks...")
    pair_mp = prep_pair("MSFT", "POOL", 600.0, 100.0)
    verdict_mp = validate_pair(pair_mp)

    path_cash = str(OUT_DIR / "p7_msft_pool_allcash.xlsx")
    generate_ma_workbook(pair_mp, verdict_mp, path_cash,
                         deal_assumptions={"pct_cash": 1.0, "pct_debt": 0.0})

    path_stock = str(OUT_DIR / "p7_msft_pool_allstock.xlsx")
    generate_ma_workbook(pair_mp, verdict_mp, path_stock,
                         deal_assumptions={"pct_cash": 0.0, "pct_debt": 0.0})

    pair_jw = prep_pair("JNJ", "WING", 160.0, 400.0)
    verdict_jw = validate_pair(pair_jw)
    path_jw = str(OUT_DIR / "p7_jnj_wing.xlsx")
    refs_jw = generate_ma_workbook(pair_jw, verdict_jw, path_jw)

    # Different market share tables for Mode B check 3
    refs_cash = load_refs_from_workbook(path_cash)
    fill_shares(path_cash, refs_cash, [0.30, 0.10, 0.25])
    fill_shares(path_jw, refs_jw, [0.45, 0.20])

    # ==================================================================
    # MODE A: all-cash vs all-stock
    # ==================================================================
    print(f"\n{'='*70}\nMODE A: all-cash vs all-stock (same deal)\n{'='*70}")
    res_a = ct.compare_files(path_cash, path_stock)

    check("mode detected as scenario", res_a.mode == "scenario", res_a.mode)

    # 1. Input diff shows ONLY the financing mix
    diff_fields = {d.field_name for d in res_a.input_diffs}
    check("input diff contains ONLY financing mix fields",
          diff_fields == {"pct_cash"},
          f"diff fields = {sorted(diff_fields)}")

    # 2. Output shift direction: reason first, then check.
    # Per marginal dollar of consideration, all-cash costs after-tax
    # foregone interest = foregone_rate*(1-tax); all-stock costs EPS
    # dilution at the acquirer's pro-forma earnings yield ~ EPS/price.
    data = res_a.data_a
    foregone = data.get("foregone_cash_rate")
    tax = data.get("tax_rate")
    standalone_eps = data.get("standalone_eps")
    acq_price = data.get("acq_price")
    cash_cost = foregone * (1 - tax)
    stock_cost = standalone_eps / acq_price
    expected = "stock" if stock_cost < cash_cost else "cash"
    y1_cash, y1_stock = res_a.ad_by_year_a[0], res_a.ad_by_year_b[0]
    actual = "stock" if y1_stock > y1_cash else "cash"
    print(f"  after-tax cash cost {cash_cost:.3%} vs stock earnings-yield "
          f"cost {stock_cost:.3%} -> expect {expected} structure better")
    print(f"  Y1 A/D: all-cash {y1_cash:+.3%} vs all-stock {y1_stock:+.3%}")
    check("output shift matches reasoned direction", expected == actual,
          f"expected {expected} better, actual {actual} better")

    # Independent first-principles recompute of both Y1 A/D values
    # (financing structure is the only difference; EBIT side identical)
    uses = data.get("total_uses")
    ebit_pre = (res_a.data_a.get("comb_ebit_year1")
                + data.get("incremental_da"))  # EBIT before inc-D&A subtraction
    inc_da = data.get("incremental_da")
    acq_shares = data.get("acq_diluted_shares")

    def y1_ad(cash_frac, stock_frac):
        pretax = (ebit_pre - inc_da
                  - uses * cash_frac * foregone)
        ni = pretax - max(0, pretax * tax)
        shares = acq_shares + (uses * stock_frac) / acq_price
        return (ni / shares) / standalone_eps - 1

    exp_cash_ad = y1_ad(1.0, 0.0)
    exp_stock_ad = y1_ad(0.0, 1.0)
    check("all-cash Y1 A/D matches first-principles recompute",
          abs(exp_cash_ad - y1_cash) < 1e-9,
          f"sheet={y1_cash!r} vs recompute={exp_cash_ad!r}")
    check("all-stock Y1 A/D matches first-principles recompute",
          abs(exp_stock_ad - y1_stock) < 1e-9,
          f"sheet={y1_stock!r} vs recompute={exp_stock_ad!r}")

    # 3. Commentary prompt discipline
    prompt_a = ct.build_scenario_comparison_prompt(res_a)
    (OUT_DIR / "p7_prompt_mode_a.txt").write_text(prompt_a, encoding="utf-8")
    check("single-variable causal exception granted (mix = one variable)",
          "EXACTLY ONE independent variable changed" in prompt_a)
    check("scope prohibition present (Mode A)",
          "regulatory approval likelihood" in prompt_a)
    check("no-speculation constraint present",
          "No speculation about deal strategy" in prompt_a)
    check("HHI treated as differing tables note",
          res_a.market_table_differs, "tables intentionally differ")

    # Multi-variable case: change premium AND tax -> exception must NOT fire
    path_multi = str(OUT_DIR / "p7_msft_pool_multi.xlsx")
    generate_ma_workbook(pair_mp, verdict_mp, path_multi,
                         deal_assumptions={"pct_cash": 1.0, "pct_debt": 0.0,
                                           "offer_premium": 0.35,
                                           "tax_rate": 0.30})
    res_multi = ct.compare_files(path_cash, path_multi)
    prompt_multi = ct.build_scenario_comparison_prompt(res_multi)
    check("multi-variable change denies causal language",
          "MULTIPLE variables changed" in prompt_multi
          and "EXACTLY ONE" not in prompt_multi)

    # ==================================================================
    # MODE B: two different deals
    # ==================================================================
    print(f"\n{'='*70}\nMODE B: different deals\n{'='*70}")
    res_b = ct.compare_files(path_cash, path_jw)

    check("mode detected as deal", res_b.mode == "deal", res_b.mode)
    check("deal labels correct",
          res_b.deal_label_a == "MSFT acquiring POOL"
          and res_b.deal_label_b == "JNJ acquiring WING",
          f"{res_b.deal_label_a} / {res_b.deal_label_b}")

    # 1. Full assumptions side by side (all 13, not just diffs)
    check("all 13 assumptions side by side",
          len(res_b.assumptions_comparison) == len(ct.ASSUMPTION_FIELDS),
          f"{len(res_b.assumptions_comparison)}")
    check("no input-diff list in deal mode", res_b.input_diffs == [])

    prompt_b = ct.build_deal_comparison_prompt(res_b)
    (OUT_DIR / "p7_prompt_mode_b.txt").write_text(prompt_b, encoding="utf-8")

    # 2. Deal A vs Deal B framing
    check("prompt uses Deal A/Deal B framing",
          "Deal A: MSFT acquiring POOL" in prompt_b
          and "Deal B: JNJ acquiring WING" in prompt_b)
    check("prompt forbids base-case delta language",
          "NOT a base case and a change" in prompt_b)
    check("both fundamentals blocks present",
          "DEAL A PROFILE" in prompt_b and "DEAL B PROFILE" in prompt_b)

    # 3. HHI: figures stated, no clearance speculation invited
    check("both deals' HHI figures in prompt",
          "Deal A: pre" in prompt_b and "Deal B: pre" in prompt_b)
    check("antitrust-clearance prohibition present (Mode B)",
          "Do NOT speculate about which deal is more likely to clear "
          "antitrust review" in prompt_b)
    check("scope prohibition present (Mode B)",
          "regulatory approval likelihood" in prompt_b)

    # Deterministic printout should not crash in either mode
    ct.print_comparison_result(res_a)
    ct.print_comparison_result(res_b)

    print(f"\n{'='*70}")
    if failures:
        print(f"PHASE 7 TEST RESULT: {len(failures)} FAILURE(S)")
        for f in failures:
            print(f"  - {f}")
    else:
        print("PHASE 7 TEST RESULT: ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
